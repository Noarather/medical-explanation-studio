"""MedLearning Question Exchange v2 shared contract.

This module is deliberately dependency-light.  JSON helpers use the standard
library; XLSX helpers import openpyxl only when called.
"""
from __future__ import annotations

from copy import deepcopy
from datetime import datetime, timezone
from hashlib import sha1
import json
import re
import unicodedata
from pathlib import Path
from typing import Any, Iterable

FORMAT_NAME = "medlearning.question-set"
SCHEMA_VERSION = 2
SECTION_TITLES = {
    "analysis": "考点解析",
    "answerBasis": "正确答案依据",
    "pitfalls": "易错点提示",
    "clinicalNotes": "临床与操作要点",
}
BLOCK_TYPES = {"paragraph", "list", "table", "callout"}
CALLOUT_TONES = {"info", "important", "warning"}
CORE_KEYS = (
    "id", "bank", "type", "subject", "system", "difficulty", "caseInfo",
    "question", "options", "answer", "year", "sourceType", "questionSource",
)
STANDARD_KEYS = set(CORE_KEYS) | {
    "explanation", "explanationBlocks", "knowledgePoint", "knowledgePoints", "studyPoints",
    "memoryCards", "memoryCardMeta",
    "briefExplanation", "tags", "suggestedTags", "mnemonic", "explanationMeta", "extensions",
    "analysis", "evidence", "knowledge", "记忆口诀", "口诀", "解析",
}


def _text(value: Any, limit: int = 100_000) -> str:
    return str(value if value is not None else "").strip().replace("\r\n", "\n").replace("\r", "\n")[:limit]


def _one_line(value: Any, limit: int = 500) -> str:
    return re.sub(r"\s+", " ", _text(value, limit)).strip()


def _plain_content(value: Any, limit: int = 100_000) -> str:
    return re.sub(r"<[^>]*>", "", re.sub(r"<br\s*/?>", "\n", _text(value, limit), flags=re.I))


def _limited_list(value: Any, max_items: int, max_length: int) -> list[str]:
    source = value if isinstance(value, list) else re.split(r"[|,，;；、/\n]+", str(value or ""))
    result: list[str] = []
    seen: set[str] = set()
    for item in source:
        normalized = _one_line(item, max_length)
        key = normalized.casefold()
        if len(normalized) < 2 or key in seen:
            continue
        seen.add(key)
        result.append(normalized)
        if len(result) >= max_items:
            break
    return result


def normalize_knowledge_points(value: Any) -> list[str]:
    return _limited_list(value, 3, 60)


def normalize_tags(value: Any) -> list[str]:
    return _limited_list(value, 3, 24)


def normalize_fill_answers(value: Any) -> list[list[str]]:
    """fill 题型答案：二维数组，每空候选非空、≤60 字、去重；任一空非法则整体返回 []。
    字符串输入（如 XLSX 答案列）先尝试 json.loads。语义与 web 版 fillAnswers 一致。"""
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except (TypeError, json.JSONDecodeError):
            return []
    if not isinstance(value, list):
        return []
    result: list[list[str]] = []
    for blank in value:
        if not isinstance(blank, list):
            return []
        candidates = [candidate for candidate in (_one_line(item, 60) for item in blank) if candidate]
        if not candidates:
            return []
        result.append(list(dict.fromkeys(candidates)))
    return result


def normalize_study_points(value) -> list[dict]:
    """Review-grade study points: <=3 items, title <=20 chars, body <=300 chars, deduped by title."""
    if not isinstance(value, list):
        return []
    seen: set[str] = set()
    out: list[dict] = []
    for item in value:
        if not isinstance(item, dict):
            continue
        title = str(item.get("title") or "").strip()[:20]
        body = str(item.get("body") or "").strip()[:300]
        if not title or not body or title in seen:
            continue
        seen.add(title)
        out.append({"title": title, "body": body})
        if len(out) >= 3:
            break
    return out


def normalize_memory_cards(value: Any, raw: dict[str, Any] | None = None) -> list[dict]:
    """Independent review cards; exact copies of legacy labels/explanations are rejected."""
    if not isinstance(value, list):
        return []
    source = raw if isinstance(raw, dict) else {}

    def comparable(item: Any) -> str:
        return re.sub(
            r"[\W_]+", "", unicodedata.normalize("NFKC", _one_line(_plain_content(item, 5000), 5000)).casefold(),
        )

    blocked_values: list[Any] = []
    for key in ("knowledgePoints", "tags", "suggestedTags"):
        candidate = source.get(key)
        blocked_values.extend(candidate if isinstance(candidate, list) else [candidate])
    blocked_values.extend((source.get("briefExplanation"), source.get("explanation"), source.get("mnemonic")))
    for point in source.get("studyPoints") if isinstance(source.get("studyPoints"), list) else []:
        if isinstance(point, dict):
            blocked_values.extend((point.get("title"), point.get("body")))
    for block in source.get("explanationBlocks") if isinstance(source.get("explanationBlocks"), list) else []:
        if not isinstance(block, dict):
            continue
        blocked_values.extend((block.get("title"), block.get("text")))
        blocked_values.extend(block.get("items") if isinstance(block.get("items"), list) else [])
        for row in block.get("rows") if isinstance(block.get("rows"), list) else []:
            blocked_values.extend(row if isinstance(row, list) else [])
    blocked = {comparable(item) for item in blocked_values if comparable(item)}
    seen: set[str] = set()
    result: list[dict] = []
    for item in value:
        if not isinstance(item, dict):
            continue
        title = _one_line(item.get("title"), 30)
        content = _text(item.get("content") or item.get("body"), 500)
        memory_cue = _one_line(item.get("memoryCue"), 160)
        contrast = _text(item.get("contrast"), 300)
        title_key, content_key = comparable(title), comparable(content)
        if not title or not content or title_key in seen or title_key in blocked or content_key in blocked:
            continue
        seen.add(title_key)
        card = {"title": title, "content": content}
        if memory_cue:
            card["memoryCue"] = memory_cue
        if contrast:
            card["contrast"] = contrast
        result.append(card)
        if len(result) >= 3:
            break
    return result


def normalize_memory_card_meta(value: Any) -> dict[str, Any]:
    source = value if isinstance(value, dict) else {}
    return {
        "schemaVersion": max(1, min(10, int(source.get("schemaVersion") or 1))),
        "generator": _one_line(source.get("generator") or "analysis-tool", 80),
        "generatedAt": _one_line(source.get("generatedAt"), 40),
    }


def _generation_trace(value: Any) -> list[dict[str, Any]]:
    result = []
    for item in value[:20] if isinstance(value, list) else []:
        if not isinstance(item, dict):
            continue
        result.append({
            "finishReason": _one_line(item.get("finishReason"), 40),
            "model": _one_line(item.get("model"), 80),
            "thinking": bool(item.get("thinking", False)),
            "elapsedMs": max(0, int(item.get("elapsedMs") or 0)),
            "promptTokens": max(0, int(item.get("promptTokens") or 0)),
            "completionTokens": max(0, int(item.get("completionTokens") or 0)),
            "reasoningTokens": max(0, int(item.get("reasoningTokens") or 0)),
            "cachedPromptTokens": max(0, int(item.get("cachedPromptTokens") or 0)),
            "cacheStatus": _one_line(item.get("cacheStatus"), 16),
            "request": max(0, int(item.get("request") or 0)),
            "compactRegeneration": bool(item.get("compactRegeneration", False)),
        })
    return result


def _block_id(block: dict[str, Any], index: int) -> str:
    existing = re.sub(r"[^a-zA-Z0-9_-]", "", _one_line(block.get("blockId"), 40))
    if existing:
        return existing
    digest = sha1(json.dumps(block, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()[:6]
    return f"b{index + 1:02d}-{digest}"


def legacy_explanation_to_blocks(value: Any) -> list[dict[str, Any]]:
    source = _plain_content(value, 50_000)
    if not source:
        return []
    chunks = [item.strip() for item in re.split(r"(?=^###\s+)", source, flags=re.M) if item.strip()]
    result: list[dict[str, Any]] = []
    for index, chunk in enumerate(chunks[:12]):
        match = re.match(r"^###\s*([^\n]+)\n?", chunk)
        title = _one_line(match.group(1), 40) if match else (SECTION_TITLES["analysis"] if index == 0 else "")
        body = chunk[match.end():].strip() if match else chunk
        section = (
            "answerBasis" if re.search(r"答案|依据", title) else
            "pitfalls" if re.search(r"易错|辨析", title) else
            "clinicalNotes" if re.search(r"临床|操作|关键", title) else "analysis"
        )
        lines = [line.strip() for line in body.splitlines() if line.strip()]
        raw = {"title": title, "body": body}
        if len(lines) >= 2 and all(re.match(r"^[-*]\s+", line) for line in lines):
            result.append({"blockId": _block_id(raw, index), "section": section, "type": "list", "title": title,
                           "items": [re.sub(r"^[-*]\s+", "", line) for line in lines[:20]]})
        else:
            result.append({"blockId": _block_id(raw, index), "section": section, "type": "paragraph", "title": title,
                           "text": body[:4000]})
    return result


def normalize_blocks(value: Any, fallback: Any = "") -> list[dict[str, Any]]:
    # Some legacy imports stored the complete structured payload as a JSON string.
    # Decode it here so the website never renders raw `explanationBlocks` JSON.
    for candidate in (value, fallback):
        if not isinstance(candidate, str):
            continue
        encoded = re.sub(r"^```(?:json)?\s*|\s*```$", "", candidate.strip(), flags=re.I | re.S)
        if not encoded.startswith(("{", "[")):
            continue
        try:
            decoded = json.loads(encoded)
        except (TypeError, ValueError, json.JSONDecodeError):
            continue
        decoded_blocks = decoded.get("explanationBlocks") if isinstance(decoded, dict) else decoded
        if isinstance(decoded_blocks, list) and decoded_blocks:
            value = decoded_blocks
            break
    if not isinstance(value, list) or not value:
        return legacy_explanation_to_blocks(fallback)
    result: list[dict[str, Any]] = []
    for index, raw in enumerate(value[:12]):
        if not isinstance(raw, dict):
            continue
        block_type = raw.get("type") if raw.get("type") in BLOCK_TYPES else "paragraph"
        section = raw.get("section") if raw.get("section") in SECTION_TITLES else "analysis"
        block: dict[str, Any] = {
            "blockId": _block_id(raw, index), "section": section, "type": block_type,
            "title": _one_line(raw.get("title"), 40),
        }
        if block_type in {"paragraph", "callout"}:
            block["text"] = _plain_content(raw.get("text", raw.get("content", "")), 4000)
            if not block["text"]:
                continue
        if block_type == "callout":
            block["tone"] = raw.get("tone") if raw.get("tone") in CALLOUT_TONES else "info"
        if block_type == "list":
            source = raw.get("items") if isinstance(raw.get("items"), list) else str(raw.get("text", "")).splitlines()
            block["items"] = [_one_line(_plain_content(item, 500), 500) for item in source if _one_line(_plain_content(item, 500), 500)][:20]
            if not block["items"]:
                continue
        if block_type == "table":
            block["columns"] = [_one_line(_plain_content(item, 80), 80) for item in (raw.get("columns") or []) if _one_line(_plain_content(item, 80), 80)][:5]
            if len(block["columns"]) < 2:
                continue
            rows: list[list[str]] = []
            for raw_row in (raw.get("rows") or [])[:20]:
                if not isinstance(raw_row, list):
                    continue
                cells = [_plain_content(cell, 500) for cell in raw_row[:len(block["columns"])]]
                cells.extend([""] * (len(block["columns"]) - len(cells)))
                if any(cells):
                    rows.append(cells)
            block["rows"] = rows
            if not rows:
                continue
        result.append(block)
    return result or legacy_explanation_to_blocks(fallback)


def normalize_evidence(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    result = []
    for raw in value[:10]:
        if not isinstance(raw, dict):
            continue
        source_page = raw.get("sourcePage", raw.get("source_page"))
        pdf_page = raw.get("pdfPage", raw.get("pdf_page"))
        result.append({
            "textbook": _one_line(raw.get("textbook") or raw.get("library") or raw.get("source_file") or raw.get("fileName"), 120),
            "version": _one_line(raw.get("version") or raw.get("textbook_version"), 60),
            "fileName": _one_line(raw.get("fileName") or raw.get("source_file"), 180),
            "sourcePage": int(source_page) if str(source_page or "").isdigit() and int(source_page) > 0 else None,
            "pdfPage": int(pdf_page) if str(pdf_page or "").isdigit() and int(pdf_page) > 0 else None,
            "quote": _text(raw.get("quote") or raw.get("text"), 2000),
            "excerpt": _text(raw.get("excerpt"), 500),
            "contextSummary": _text(raw.get("contextSummary") or raw.get("context_summary"), 500),
            "highlights": [_one_line(item, 120) for item in (raw.get("highlights") or []) if _one_line(item, 120)][:8],
            "supportedClues": [_one_line(item, 120) for item in (raw.get("supportedClues") or raw.get("supported_clues") or []) if _one_line(item, 120)][:8],
            "score": float(raw.get("score") or 0),
            "extractionMethod": _one_line(raw.get("extractionMethod") or raw.get("extraction_method") or "text", 20),
        })
    return result


def blocks_to_markdown(blocks: Any, evidence: Any = None) -> str:
    parts: list[str] = []
    previous = ""
    for block in normalize_blocks(blocks):
        if block["section"] != previous:
            parts.append(f"### {SECTION_TITLES[block['section']]}")
            previous = block["section"]
        if block.get("title") and block["title"] != SECTION_TITLES[block["section"]]:
            parts.append(f"**{block['title']}**")
        if block["type"] == "paragraph":
            parts.append(block["text"])
        elif block["type"] == "callout":
            prefix = "注意：" if block.get("tone") == "warning" else "重点：" if block.get("tone") == "important" else ""
            parts.append(f"> {prefix}{block['text']}")
        elif block["type"] == "list":
            parts.append("\n".join(f"- {item}" for item in block["items"]))
        elif block["type"] == "table":
            esc = lambda v: _text(v, 500).replace("|", "\\|").replace("\n", "<br>")
            parts.append(
                f"| {' | '.join(map(esc, block['columns']))} |\n"
                f"| {' | '.join('---' for _ in block['columns'])} |\n" +
                "\n".join(f"| {' | '.join(map(esc, row))} |" for row in block["rows"])
            )
    sources = normalize_evidence(evidence or [])
    if sources:
        parts.append("### 教材出处")
        parts.append("\n".join(
            f"- 《{item['textbook'] or item['fileName'] or '教材'}》"
            f"{'（' + item['version'] + '）' if item['version'] else ''} · "
            f"{'课本第' + str(item['sourcePage']) + '页' if item['sourcePage'] else '课本前置页'}"
            for item in sources
        ))
    return "\n\n".join(part for part in parts if part)


def normalize_question_v2(raw: dict[str, Any]) -> dict[str, Any]:
    raw = deepcopy(raw or {})
    evidence = normalize_evidence((raw.get("explanationMeta") or {}).get("evidence") or raw.get("evidence"))
    blocks = normalize_blocks(raw.get("explanationBlocks"), raw.get("explanation") or raw.get("analysis") or raw.get("解析"))
    extensions = deepcopy(raw.get("extensions")) if isinstance(raw.get("extensions"), dict) else {}
    for key, value in raw.items():
        if key not in STANDARD_KEYS:
            extensions[key] = value
    meta = raw.get("explanationMeta") if isinstance(raw.get("explanationMeta"), dict) else {}
    result = {key: raw.get(key) for key in CORE_KEYS}
    mode = _one_line(meta.get("mode"), 40)
    inferred_grade = "C" if mode == "general_knowledge" else "B" if mode == "textbook_reasoning" else "A" if mode == "textbook" else ""
    evidence_grade = _one_line(meta.get("evidenceGrade") or inferred_grade, 1).upper()
    if evidence_grade not in {"A", "B", "C"}:
        evidence_grade = ""
    reasoning_type = _one_line(meta.get("reasoningType"), 24)
    if reasoning_type not in {"direct_answer", "case_chain"}:
        reasoning_type = ""
    result.update({
        "answer": normalize_fill_answers(raw.get("answer")) if _one_line(raw.get("type"), 20) == "fill" else raw.get("answer"),
        "questionSource": unicodedata.normalize("NFKC", _one_line(raw.get("questionSource") or raw.get("题库来源"), 40)),
        "explanationBlocks": blocks,
        "explanation": blocks_to_markdown(blocks, evidence),
        "briefExplanation": _one_line(raw.get("briefExplanation") or raw.get("shortExplanation") or raw.get("简短解析"), 160),
        "knowledgePoints": normalize_knowledge_points(raw.get("knowledgePoints") or raw.get("knowledgePoint") or raw.get("knowledge")),
        "studyPoints": normalize_study_points(raw.get("studyPoints")),
        "memoryCards": normalize_memory_cards(raw.get("memoryCards"), raw),
        "memoryCardMeta": normalize_memory_card_meta(raw.get("memoryCardMeta")),
        "tags": normalize_tags(raw.get("tags")),
        "suggestedTags": normalize_tags(raw.get("suggestedTags")),
        "mnemonic": _text(raw.get("mnemonic") or raw.get("记忆口诀") or raw.get("口诀"), 1000),
        "explanationMeta": {
            "mode": mode, "model": _one_line(meta.get("model"), 80),
            "evidenceGrade": evidence_grade, "reasoningType": reasoning_type,
            "thinking": bool(meta.get("thinking", False)),
            "modelFallback": bool(meta.get("modelFallback", False)),
            "complexityScore": max(0, min(20, int(meta.get("complexityScore") or 0))),
            "routeReasons": _limited_list(meta.get("routeReasons"), 12, 80),
            "requestCount": max(0, min(20, int(meta.get("requestCount") or 0))),
            "generationTrace": _generation_trace(meta.get("generationTrace")),
            "score": float(meta.get("score") or 0),
            "generatedAt": _one_line(meta.get("generatedAt") or meta.get("generated_at"), 40),
            "reviewedAt": _one_line(meta.get("reviewedAt") or meta.get("reviewed_at"), 40),
            "reviewWarnings": _limited_list(meta.get("reviewWarnings"), 10, 200),
            "evidence": evidence,
        },
        "extensions": extensions,
    })
    result["knowledgePoint"] = "、".join(result["knowledgePoints"])
    return result


def make_envelope(questions: Iterable[dict[str, Any]], *, name: str = "", bank: str = "", source: str = "unknown",
                  exported_at: str = "", tag_catalog: Iterable[dict[str, Any]] = ()) -> dict[str, Any]:
    normalized = [normalize_question_v2(item) for item in questions]
    bank = _one_line(bank or (normalized[0].get("bank") if normalized else ""), 20)
    catalog = []
    for item in tag_catalog:
        entry = {
            "subject": _one_line(item.get("subject"), 60), "label": _one_line(item.get("label") or item.get("tag"), 24),
            "aliases": _limited_list(item.get("aliases"), 20, 24),
            "status": item.get("status") if item.get("status") in {"active", "candidate", "merged"} else "active",
            "usageCount": max(0, int(item.get("usageCount") or item.get("count") or 0)),
            "mergedInto": _one_line(item.get("mergedInto"), 24),
        }
        if entry["subject"] and entry["label"]:
            catalog.append(entry)
    return {
        "schema": FORMAT_NAME, "schemaVersion": SCHEMA_VERSION,
        "meta": {"name": _one_line(name or f"{bank or 'medical'}-questions", 120), "bank": bank,
                 "source": _one_line(source, 40),
                 "exportedAt": exported_at or datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
                 "questionCount": len(normalized)},
        "tagCatalog": catalog, "questions": normalized,
    }


def is_v2(value: Any) -> bool:
    return isinstance(value, dict) and value.get("schema") == FORMAT_NAME and int(value.get("schemaVersion") or 0) == 2 and isinstance(value.get("questions"), list)


def write_json_v2(path: str | Path, envelope: dict[str, Any]) -> None:
    Path(path).write_text(json.dumps(envelope, ensure_ascii=False, indent=2), encoding="utf-8")


QUESTION_HEADERS = [
    "ID", "题库", "题库来源", "题型", "学科", "章节", "难度", "病例背景", "题干", "选项A", "选项B", "选项C", "选项D", "选项E",
    "答案", "知识点1", "知识点2", "知识点3", "题目标签1", "题目标签2", "题目标签3", "候选标签1", "候选标签2", "候选标签3",
    "一句话简析", "记忆口诀", "年份", "解析纯文本", "来源类型", "证据等级", "推理类型", "考点JSON", "扩展字段JSON",
]

# XML 1.0 forbids these control characters inside worksheet cells. They can
# appear after copying textbook text from PDF/Word even though JSON accepts
# them. Keep tabs/newlines/carriage returns, which Excel supports.
_XLSX_FORBIDDEN_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F\uFFFE\uFFFF]")


def _xlsx_cell(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    return _XLSX_FORBIDDEN_CHARACTERS.sub("", value)[:32767]


def _xlsx_row(values: Iterable[Any]) -> list[Any]:
    return [_xlsx_cell(value) for value in values]


def write_xlsx_v2(path: str | Path, envelope: dict[str, Any]) -> None:
    from openpyxl import Workbook
    value = make_envelope(envelope.get("questions", []), name=(envelope.get("meta") or {}).get("name", ""),
                          bank=(envelope.get("meta") or {}).get("bank", ""), source=(envelope.get("meta") or {}).get("source", "unknown"),
                          exported_at=(envelope.get("meta") or {}).get("exportedAt", ""), tag_catalog=envelope.get("tagCatalog", []))
    wb = Workbook(); meta_ws = wb.active; meta_ws.title = "_meta"
    meta_ws.append(_xlsx_row(["key", "value"]))
    for key, item in {"schema": value["schema"], "schemaVersion": 2, **value["meta"]}.items(): meta_ws.append(_xlsx_row([key, item]))
    meta_ws.sheet_state = "hidden"
    q_ws = wb.create_sheet("题目"); q_ws.append(_xlsx_row(QUESTION_HEADERS))
    b_ws = wb.create_sheet("解析区块"); b_headers = ["题目ID", "区块序号", "区块ID", "所属部分", "类型", "标题", "提示样式", "行序号", "正文"] + [v for i in range(1, 6) for v in (f"列名{i}", f"单元格{i}")]; b_ws.append(_xlsx_row(b_headers))
    e_ws = wb.create_sheet("教材证据"); e_headers = ["题目ID", "序号", "教材名称", "版本", "文件名", "课本页码", "PDF页码", "相似度", "提取方式", "关键原句", "重点片段", "支持线索", "AI整理上下文", "教材原文"]; e_ws.append(_xlsx_row(e_headers))
    m_ws = wb.create_sheet("背诵知识卡"); m_ws.append(_xlsx_row(["题目ID", "序号", "标题", "独立背诵正文", "记忆提示", "对比辨析", "生成器", "生成时间", "格式版本"]))
    t_ws = wb.create_sheet("标签目录"); t_ws.append(_xlsx_row(["学科", "标签", "别名", "状态", "使用次数", "合并至"]))
    for question in value["questions"]:
        option_values = [re.sub(r"^[A-E][.、．:]\s*", "", _one_line(item)) for item in (question.get("options") or [])][:5]
        option_values.extend([""] * (5 - len(option_values)))
        kp = question["knowledgePoints"] + [""] * 3; tg = question["tags"] + [""] * 3; st = question["suggestedTags"] + [""] * 3
        answer_cell = question.get("answer")
        if isinstance(answer_cell, list):
            answer_cell = json.dumps(answer_cell, ensure_ascii=False)
        q_ws.append(_xlsx_row([question.get("id"), question.get("bank"), question.get("questionSource"), question.get("type"), question.get("subject"), question.get("system"), question.get("difficulty"), question.get("caseInfo"), question.get("question"), *option_values,
                     answer_cell, *kp[:3], *tg[:3], *st[:3], question.get("briefExplanation"), question.get("mnemonic"), question.get("year"), question.get("explanation"), question.get("sourceType"), question["explanationMeta"].get("evidenceGrade"), question["explanationMeta"].get("reasoningType"), json.dumps(question.get("studyPoints") or [], ensure_ascii=False), json.dumps(question.get("extensions") or {}, ensure_ascii=False)]))
        for block_index, block in enumerate(question["explanationBlocks"], 1):
            base = [question.get("id"), block_index, block.get("blockId"), block.get("section"), block.get("type"), block.get("title", ""), block.get("tone", "")]
            if block["type"] == "table":
                for row_index, cells in enumerate(block["rows"], 1):
                    data = base + [row_index, ""]
                    for i in range(5): data += [block["columns"][i] if i < len(block["columns"]) else "", cells[i] if i < len(cells) else ""]
                    b_ws.append(_xlsx_row(data))
            elif block["type"] == "list":
                for row_index, item in enumerate(block["items"], 1): b_ws.append(_xlsx_row(base + [row_index, item] + [""] * 10))
            else: b_ws.append(_xlsx_row(base + [1, block.get("text", "")] + [""] * 10))
        for index, item in enumerate(question["explanationMeta"]["evidence"], 1):
            e_ws.append(_xlsx_row([question.get("id"), index, item["textbook"], item["version"], item["fileName"], item["sourcePage"], item["pdfPage"], item["score"], item["extractionMethod"], item.get("excerpt", ""), "、".join(item.get("highlights") or []), "、".join(item.get("supportedClues") or []), item.get("contextSummary", ""), item["quote"]]))
        for index, item in enumerate(question.get("memoryCards") or [], 1):
            meta = question.get("memoryCardMeta") or {}
            m_ws.append(_xlsx_row([question.get("id"), index, item.get("title"), item.get("content"), item.get("memoryCue", ""), item.get("contrast", ""), meta.get("generator", "analysis-tool"), meta.get("generatedAt", ""), meta.get("schemaVersion", 1)]))
    for item in value["tagCatalog"]: t_ws.append(_xlsx_row([item["subject"], item["label"], "、".join(item["aliases"]), item["status"], item["usageCount"], item["mergedInto"]]))
    wb.save(path)


def read_xlsx_v2(path: str | Path) -> dict[str, Any] | None:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    if "_meta" not in wb.sheetnames or "题目" not in wb.sheetnames:
        wb.close(); return None
    def rows(name: str) -> list[dict[str, Any]]:
        ws = wb[name]; values = ws.iter_rows(values_only=True); headers = [str(item or "") for item in next(values, [])]
        return [dict(zip(headers, row)) for row in values]
    meta = {str(row.get("key") or ""): row.get("value") for row in rows("_meta")}
    if meta.get("schema") != FORMAT_NAME or int(meta.get("schemaVersion") or 0) != 2:
        wb.close(); return None
    block_groups: dict[tuple[str, int], list[dict[str, Any]]] = {}
    for row in rows("解析区块") if "解析区块" in wb.sheetnames else []:
        block_groups.setdefault((_one_line(row.get("题目ID")), int(row.get("区块序号") or 1)), []).append(row)
    blocks_by_id: dict[str, list[dict[str, Any]]] = {}
    for (question_id, _), group in sorted(block_groups.items()):
        first = group[0]; block_type = _one_line(first.get("类型")); block = {"blockId": first.get("区块ID"), "section": first.get("所属部分"), "type": block_type, "title": first.get("标题") or ""}
        group.sort(key=lambda row: int(row.get("行序号") or 0))
        if block_type == "table":
            columns = [_one_line(first.get(f"列名{i}"), 80) for i in range(1, 6)]; columns = [item for item in columns if item]
            block.update(columns=columns, rows=[[_text(row.get(f"单元格{i}"), 500) for i in range(1, len(columns) + 1)] for row in group])
        elif block_type == "list": block["items"] = [_one_line(row.get("正文"), 500) for row in group if _one_line(row.get("正文"), 500)]
        else:
            block["text"] = _text(first.get("正文"), 4000)
            if block_type == "callout": block["tone"] = first.get("提示样式") or "info"
        blocks_by_id.setdefault(question_id, []).append(block)
    evidence_by_id: dict[str, list[dict[str, Any]]] = {}
    for row in rows("教材证据") if "教材证据" in wb.sheetnames else []:
        evidence_by_id.setdefault(_one_line(row.get("题目ID")), []).append({"textbook": row.get("教材名称"), "version": row.get("版本"), "fileName": row.get("文件名"), "sourcePage": row.get("课本页码"), "pdfPage": row.get("PDF页码"), "score": row.get("相似度"), "extractionMethod": row.get("提取方式"), "excerpt": row.get("关键原句"), "highlights": [item for item in str(row.get("重点片段") or "").split("、") if item], "supportedClues": [item for item in str(row.get("支持线索") or "").split("、") if item], "contextSummary": row.get("AI整理上下文"), "quote": row.get("教材原文") or row.get("完整上下文") or row.get("证据摘要")})
    memory_cards_by_id: dict[str, list[dict[str, Any]]] = {}
    memory_meta_by_id: dict[str, dict[str, Any]] = {}
    for row in rows("背诵知识卡") if "背诵知识卡" in wb.sheetnames else []:
        question_id = _one_line(row.get("题目ID"))
        if not question_id:
            continue
        memory_cards_by_id.setdefault(question_id, []).append({
            "order": int(row.get("序号") or 1), "title": row.get("标题"),
            "content": row.get("独立背诵正文"), "memoryCue": row.get("记忆提示"),
            "contrast": row.get("对比辨析"),
        })
        memory_meta_by_id.setdefault(question_id, {
            "schemaVersion": int(row.get("格式版本") or 1), "generator": row.get("生成器"),
            "generatedAt": row.get("生成时间"),
        })
    questions = []
    for row in rows("题目"):
        question_id = _one_line(row.get("ID")); options = [f"{chr(64+i)}. {_one_line(row.get(f'选项{chr(64+i)}'))}" for i in range(1, 6) if _one_line(row.get(f"选项{chr(64+i)}"))]
        try: extensions = json.loads(row.get("扩展字段JSON") or "{}")
        except (TypeError, json.JSONDecodeError): extensions = {"importWarning": "扩展字段JSON无法解析"}
        try: study_points = json.loads(row.get("考点JSON") or "[]")
        except (TypeError, json.JSONDecodeError): study_points = []
        if not isinstance(study_points, list): study_points = []
        answer_value = row.get("答案")
        if _one_line(row.get("题型"), 20) == "fill" and isinstance(answer_value, str) and answer_value.strip().startswith("["):
            try: answer_value = json.loads(answer_value)
            except json.JSONDecodeError: answer_value = []
        memory_cards = sorted(memory_cards_by_id.get(question_id, []), key=lambda item: item.pop("order", 1))
        questions.append(normalize_question_v2({"id": question_id, "bank": row.get("题库") or meta.get("bank"), "questionSource": row.get("题库来源"), "type": row.get("题型"), "subject": row.get("学科"), "system": row.get("章节"), "difficulty": row.get("难度"), "caseInfo": row.get("病例背景"), "question": row.get("题干"), "options": options, "answer": answer_value, "knowledgePoints": [row.get(f"知识点{i}") for i in range(1, 4)], "tags": [row.get(f"题目标签{i}") for i in range(1, 4)], "suggestedTags": [row.get(f"候选标签{i}") for i in range(1, 4)], "briefExplanation": row.get("一句话简析"), "mnemonic": row.get("记忆口诀"), "year": row.get("年份"), "explanation": row.get("解析纯文本"), "explanationBlocks": blocks_by_id.get(question_id, []), "explanationMeta": {"evidence": evidence_by_id.get(question_id, []), "evidenceGrade": row.get("证据等级"), "reasoningType": row.get("推理类型")}, "sourceType": row.get("来源类型"), "studyPoints": study_points, "memoryCards": memory_cards, "memoryCardMeta": memory_meta_by_id.get(question_id), "extensions": extensions}))
    catalog = [{"subject": row.get("学科"), "label": row.get("标签"), "aliases": row.get("别名"), "status": row.get("状态"), "usageCount": row.get("使用次数"), "mergedInto": row.get("合并至")} for row in (rows("标签目录") if "标签目录" in wb.sheetnames else [])]
    wb.close()
    return make_envelope(questions, name=meta.get("name") or "", bank=meta.get("bank") or "", source=meta.get("source") or "unknown", exported_at=meta.get("exportedAt") or "", tag_catalog=catalog)
