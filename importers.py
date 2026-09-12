"""Question JSON/Excel import validation while preserving unknown source fields."""

from __future__ import annotations

import json
import copy
import hashlib
import re
import unicodedata
from pathlib import Path
from typing import Any, Callable
from question_format_v2 import (
    is_v2, normalize_blocks, normalize_question_v2, normalize_tags,
    normalize_knowledge_points, normalize_fill_answers, read_xlsx_v2,
)


STANDARD_FIELDS = [
    "id", "bank", "type", "subject", "system", "difficulty", "caseInfo",
    "question", "options", "answer", "explanation", "explanationBlocks", "knowledgePoint", "knowledgePoints",
    "briefExplanation", "studyPoints", "memoryCards", "memoryCardMeta", "tags", "suggestedTags", "mnemonic", "explanationMeta", "extensions", "sourceType", "questionSource", "year",
]
MAPPING_FIELDS = STANDARD_FIELDS + ["optionA", "optionB", "optionC", "optionD", "optionE"]
REQUIRED_FIELDS = ("id", "subject", "question", "answer")


def question_prompt(question: dict) -> str:
    option_text = "\n".join(str(item) for item in (question.get("options") or []))
    question_type = normalize_question_type(question.get("type"), question.get("answer"))
    extensions = question.get("extensions") if isinstance(question.get("extensions"), dict) else {}
    original_type = str(extensions.get("originalQuestionType") or "").strip()
    type_text = exam_question_type_label(original_type, question_type)
    if not type_text:
        type_text = {
        "A1": "A1 基础单选", "A2": "A2 病例单选", "A3": "A3/A4 病例组单选",
        "multiple": "多选题", "judge": "判断题",
        }.get(question_type, str(question_type or "未知"))
    return "\n".join(
        part for part in [
            f"题型：{type_text}",
            str(question.get("caseInfo") or "").strip(),
            str(question.get("question") or question.get("title") or "").strip(),
            option_text,
            f"正确答案：{question.get('answer') or question.get('correctAnswer') or ''}",
        ] if part
    )


class ImportValidationError(ValueError):
    def __init__(self, errors: list[str]):
        super().__init__("\n".join(errors))
        self.errors = errors


def _options(value: Any) -> list[str]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    if not text:
        return []
    if text.startswith("["):
        try:
            parsed = json.loads(text)
            if isinstance(parsed, list):
                return [str(item).strip() for item in parsed if str(item).strip()]
        except json.JSONDecodeError:
            pass
    separator = "|" if "|" in text else "\n"
    return [item.strip() for item in text.split(separator) if item.strip()]


def _knowledge_points(value: Any) -> list[str]:
    return normalize_knowledge_points(value)


def normalize_question_type(value: Any, answer: Any = "") -> str:
    """Normalize common Chinese/export aliases while preserving unknown values for validation."""
    original = str(value or "").strip()
    raw = re.sub(r"\s+", "", original).upper()
    if raw == "FILL" or "填空" in raw:
        return "fill"
    if raw in {"A1", "A2", "A3"}:
        return raw
    if raw.lower() in {"multiple", "judge"}:
        return raw.lower()
    if "判断" in raw:
        return "judge"
    answer_letters = list(dict.fromkeys(re.findall(r"[A-E]", str(answer or "").upper())))
    if raw.startswith("X") or "多选" in raw or len(answer_letters) > 1:
        return "multiple"
    if "A3" in raw or "A4" in raw:
        return "A3"
    if "A2" in raw:
        return "A2"
    if (
        "A1" in raw or raw in {"B", "B型", "B型题", "单选", "单选题", "选择题"}
        or "单项选择" in raw
    ):
        return "A1"
    return original or ("multiple" if len(answer_letters) > 1 else "A1")


def exam_question_type_label(original_type: Any, normalized_type: Any = "") -> str:
    """Return a concise learner-facing label while keeping standard answer behavior."""
    raw = re.sub(r"\s+", "", str(original_type or "")).upper()
    normalized = normalize_question_type(normalized_type or raw)
    if raw.startswith("B"):
        return "B 型题 · 配伍题"
    if raw.startswith("X") or normalized == "multiple":
        return "X 型题 · 多项选择"
    if "A3" in raw and "A4" in raw:
        return "A3/A4 型题 · 病例组最佳选择"
    if "A4" in raw:
        return "A4 型题 · 病例组递进题"
    if "A3" in raw or normalized == "A3":
        return "A3 型题 · 病例组单选"
    if "A2" in raw or normalized == "A2":
        return "A2 型题 · 病例摘要单选"
    if "A1" in raw or normalized == "A1":
        return "A1 型题 · 单句最佳选择"
    if normalized == "judge":
        return "判断题"
    return ""


def extract_question_text(path: str | Path) -> str:
    """Read raw question source from plain text, Markdown or Word."""
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix in {".txt", ".md"}:
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                return source.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError("无法识别文本文件编码")
    if suffix == ".docx":
        from docx import Document
        document = Document(source)
        lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                value = "\t".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if value:
                    lines.append(value)
        return "\n".join(lines)
    raise ValueError("原始题目文件仅支持 TXT、Markdown 和 DOCX")


_COLLECTION_FILE_RE = re.compile(r"[_\-\s]*题目集合\s*$", re.I)
_QUESTION_HEADING_RE = re.compile(r"^(?P<chapter>.+?)\s*[-–—]\s*第\s*(?P<number>\d+)\s*题\s*$")
_META_RE = re.compile(
    r"题型\s*[：:]\s*(?P<type>.*?)\s*(?:\||｜)\s*"
    r"题号\s*[：:]\s*(?P<number>.*?)\s*(?:\||｜)\s*"
    r"文件\s*[：:]\s*(?P<marker>.+?)\s*$"
)
_DOCX_FIELDS = {
    "题目": "question",
    "题干": "question",
    "选项": "options",
    "正确答案": "answer",
    "答案": "answer",
    "解析": "explanation",
    "答案解析": "explanation",
    "考点讨论": "discussion",
    "考点": "discussion",
}


def is_medical_question_collection_docx(path: str | Path) -> bool:
    """Cheaply identify the exported ``学科_题目集合.docx`` family.

    The full Word package is intentionally not opened here. Large collections contain
    tens of thousands of paragraphs, so structural validation happens once in the
    background parser instead of once during file selection and again during import.
    """
    source = Path(path)
    return source.suffix.lower() == ".docx" and bool(_COLLECTION_FILE_RE.search(source.stem))


def _collection_subject(path: Path, document_title: str, default_subject: str) -> str:
    if default_subject.strip():
        return default_subject.strip()
    title = re.sub(r"\s*[-–—]\s*题目集合\s*$", "", document_title).strip()
    if title and title != document_title:
        return title
    return _COLLECTION_FILE_RE.sub("", path.stem).strip(" _-")


def _normalize_collection_type(raw_type: str, answer: str) -> str:
    normalized = normalize_question_type(raw_type, answer)
    # B-shaped matching questions in these exports have already been expanded into
    # independent stems with their own options and answer, so they are valid A1 rows.
    return normalized if normalized in {"A1", "A2", "A3", "multiple", "judge"} else "A1"


def _normalize_collection_answer(value: str) -> str:
    text = re.sub(r"\s+", "", str(value or "")).upper()
    letters = list(dict.fromkeys(re.findall(r"[A-E]", text)))
    if letters:
        return "".join(letters)
    if text in {"对", "正确", "√"}:
        return "正确"
    if text in {"错", "错误", "×", "X"}:
        return "错误"
    return str(value or "").strip().rstrip("。.")


def _split_collection_options(lines: list[str]) -> list[str]:
    result: list[str] = []
    for source in lines:
        text = re.sub(r"\s+", " ", str(source or "")).strip()
        if not text:
            continue
        matches = list(re.finditer(r"(?:(?<=^)|(?<=\s))([A-E])\s*[.．、：:]\s*", text, re.I))
        if len(matches) > 1:
            for index, match in enumerate(matches):
                body = text[match.end():(matches[index + 1].start() if index + 1 < len(matches) else len(text))].strip()
                if body:
                    result.append(f"{match.group(1).upper()}. {body}")
            continue
        match = re.match(r"^([A-E])\s*[.．、：:]?\s*(.+)$", text, re.I)
        if match:
            result.append(f"{match.group(1).upper()}. {match.group(2).strip()}")
        else:
            label = chr(65 + len(result)) if len(result) < 5 else ""
            result.append(f"{label}. {text}" if label else text)
    return result


def _brief_from_source(*values: str) -> str:
    text = re.sub(r"\s+", " ", " ".join(value for value in values if value)).strip()
    if not text:
        return ""
    sentence = re.split(r"(?<=[。！？!?])\s*", text, maxsplit=1)[0].strip()
    return sentence[:160]


def parse_medical_question_collection_docx(
    path: str | Path,
    default_subject: str = "",
    progress: Callable[[int, int], None] | None = None,
) -> tuple[list[dict], dict]:
    """Parse styled medical collection DOCX files without flattening the whole book.

    Heading 1 supplies the real chapter, Heading 2 starts a question, and Heading 3
    switches the destination field. Directory paragraphs are skipped. This keeps
    chapter boundaries and source traceability while avoiding a huge intermediate
    plain-text buffer and an unnecessary model call.
    """
    source = Path(path)
    if source.suffix.lower() != ".docx":
        raise ValueError("结构化题目集只支持 DOCX 文件")
    from docx import Document

    document = Document(source)
    paragraphs = document.paragraphs
    first_text = next((item.text.strip() for item in paragraphs if item.text.strip()), "")
    subject = _collection_subject(source, first_text, default_subject)
    if not subject:
        raise ValueError("无法从文件名或标题识别学科")

    rows: list[dict] = []
    issues: list[str] = []
    chapter = ""
    in_directory = False
    current: dict[str, Any] | None = None
    active_field = ""
    recognized_markers = 0

    def finish_current() -> None:
        nonlocal current
        if not current:
            return
        question_text = "\n".join(current.pop("question_parts", [])).strip()
        options = _split_collection_options(current.pop("option_parts", []))
        answer = _normalize_collection_answer(" ".join(current.pop("answer_parts", [])))
        explanation = "\n\n".join(current.pop("explanation_parts", [])).strip()
        discussion = "\n\n".join(current.pop("discussion_parts", [])).strip()
        raw_type = str(current.pop("raw_type", "")).strip()
        source_marker = str(current.pop("source_marker", "")).strip()
        source_number = str(current.pop("source_number", "")).strip()
        heading_number = str(current.pop("heading_number", "")).strip()
        digest_source = "|".join((subject, source_marker, current.get("system", ""), question_text, answer))
        question_id = f"school_ykb_{hashlib.sha1(digest_source.encode('utf-8')).hexdigest()[:16]}"
        blocks: list[dict[str, Any]] = []
        if explanation:
            blocks.append({"section": "answerBasis", "type": "paragraph", "title": "原题解析", "text": explanation})
        if discussion:
            blocks.append({"section": "analysis", "type": "paragraph", "title": "考点讨论", "text": discussion})
        combined_explanation = "\n\n".join(value for value in (explanation, discussion) if value)
        row = {
            "id": question_id,
            "bank": "school",
            "type": _normalize_collection_type(raw_type, answer),
            "subject": subject,
            "system": str(current.get("system") or chapter).strip(),
            "difficulty": "medium",
            "caseInfo": "",
            "question": question_text,
            "options": options,
            "answer": answer,
            "explanation": combined_explanation,
            "explanationBlocks": blocks,
            "briefExplanation": _brief_from_source(discussion, explanation),
            "knowledgePoints": [],
            "tags": [],
            "suggestedTags": [],
            "mnemonic": "",
            "explanationMeta": {},
            "extensions": {
                "importSource": "医考帮题目集 DOCX",
                "sourceDocument": source.name,
                "sourceQuestionNumber": source_number or heading_number,
                "sourceFileMarker": source_marker,
                "originalQuestionType": raw_type,
                "sourceHeading": str(current.get("source_heading") or ""),
                "examQuestionTypeLabel": exam_question_type_label(
                    raw_type, _normalize_collection_type(raw_type, answer),
                ),
            },
            "sourceType": "yikaobang-docx",
            "year": None,
        }
        row_number = len(rows) + 1
        if re.sub(r"\s+", "", raw_type).upper().startswith("B"):
            row["extensions"]["questionGroupId"] = "ykb_b_" + hashlib.sha1(
                "|".join((subject, row["system"], source_number or heading_number)).encode("utf-8")
            ).hexdigest()[:12]
        if not question_text:
            issues.append(f"第 {row_number} 题缺少题干")
        if len(options) < 2 and row["type"] != "judge":
            issues.append(f"第 {row_number} 题选项不完整")
        if not answer:
            issues.append(f"第 {row_number} 题缺少答案")
        rows.append(row)
        current = None

    total = max(1, len(paragraphs))
    for index, paragraph in enumerate(paragraphs, start=1):
        if progress and (index == 1 or index % 250 == 0 or index == total):
            progress(index, total)
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        if style == "Heading 1":
            if text == "目录概览":
                in_directory = True
                continue
            in_directory = False
            chapter = text
            continue
        if in_directory:
            continue
        question_match = _QUESTION_HEADING_RE.match(text) if style == "Heading 2" else None
        if question_match:
            finish_current()
            current = {
                "system": chapter or question_match.group("chapter").strip(),
                "source_heading": text,
                "heading_number": question_match.group("number"),
                "question_parts": [], "option_parts": [], "answer_parts": [],
                "explanation_parts": [], "discussion_parts": [],
            }
            active_field = ""
            continue
        if current is None:
            continue
        meta_match = _META_RE.match(text)
        if meta_match and not active_field:
            current["raw_type"] = meta_match.group("type").strip()
            current["source_number"] = meta_match.group("number").strip()
            current["source_marker"] = meta_match.group("marker").strip()
            continue
        if style == "Heading 3":
            marker = re.sub(r"[：:]\s*$", "", text).strip()
            active_field = _DOCX_FIELDS.get(marker, "")
            if active_field:
                recognized_markers += 1
            continue
        if re.fullmatch(r"[─━\-_=]{10,}", text):
            continue
        target = {
            "question": "question_parts", "options": "option_parts", "answer": "answer_parts",
            "explanation": "explanation_parts", "discussion": "discussion_parts",
        }.get(active_field)
        if target:
            current[target].append(text)
    finish_current()
    if progress:
        progress(total, total)
    if not rows or recognized_markers < len(rows) * 3:
        raise ValueError("该 Word 文件不是可识别的“学科 - 题目集合”格式")
    return rows, {
        "subject": subject,
        "questionCount": len(rows),
        "chapterCount": len({str(row.get("system") or "") for row in rows}),
        "issues": issues,
        "source": str(source.resolve()),
    }


def _stable_formatted_id(question_text: str, index: int) -> str:
    digest = hashlib.sha1(question_text.encode("utf-8")).hexdigest()[:12]
    return f"formatted_{digest}_{index:03d}"


def parse_common_question_text(text: str, default_subject: str = "") -> tuple[list[dict], list[str]]:
    """Parse common numbered question blocks without calling a model."""
    normalized = str(text).replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        return [], ["没有可整理的原始文本"]
    starts = list(re.finditer(r"(?m)^\s*(?:第\s*)?(\d{1,6})\s*(?:题|[\.．、\)])\s*", normalized))
    if not starts:
        blocks = [normalized]
    else:
        blocks = [
            normalized[match.end(): (starts[index + 1].start() if index + 1 < len(starts) else len(normalized))].strip()
            for index, match in enumerate(starts)
        ]
    rows: list[dict] = []
    warnings: list[str] = []
    option_pattern = re.compile(r"(?m)^\s*([A-E])\s*[\.．、\)）:：]\s*(.+?)\s*$")
    answer_pattern = re.compile(r"(?im)^\s*(?:参考)?答案\s*[：:]\s*([^\n]+)")
    explanation_pattern = re.compile(r"(?is)^\s*(?:解析|答案解析)\s*[：:]\s*(.+)$")
    knowledge_pattern = re.compile(r"(?im)^\s*(?:knowledgePoint|知识点)\s*[:：]\s*([^\n]+)")
    for block_index, block in enumerate(blocks, start=1):
        if not block:
            continue
        answer_match = answer_pattern.search(block)
        answer = answer_match.group(1).strip() if answer_match else ""
        answer = re.sub(r"[^A-E对错正确错误√×]", "", answer, flags=re.I).upper()
        options = [f"{match.group(1).upper()}. {match.group(2).strip()}" for match in option_pattern.finditer(block)]
        first_option = option_pattern.search(block)
        end_question = first_option.start() if first_option else (answer_match.start() if answer_match else len(block))
        question_text = block[:end_question].strip()
        question_text = re.sub(r"(?m)^\s*(?:题型|学科)\s*[：:].*$", "", question_text).strip()
        question_text = re.sub(r"(?im)^\s*(?:knowledgePoint|知识点)\s*[:：].*$", "", question_text).strip()
        knowledge_match = knowledge_pattern.search(block)
        knowledge_point = knowledge_match.group(1).strip() if knowledge_match else ""
        explanation = ""
        explanation_match = explanation_pattern.search(block[answer_match.end():] if answer_match else "")
        if explanation_match:
            explanation = explanation_match.group(1).strip()
            explanation = knowledge_pattern.sub("", explanation).strip()
        question_type = "multiple" if len(re.findall(r"[A-E]", answer)) > 1 else "A1"
        if answer in {"对", "错", "正确", "错误", "√", "×"}:
            question_type = "judge"
        row = {
            "id": _stable_formatted_id(question_text or block, block_index),
            "bank": "", "type": question_type, "subject": default_subject.strip(),
            "system": "", "difficulty": "medium", "caseInfo": "",
            "question": question_text, "options": options, "answer": answer,
            "explanation": explanation, "knowledgePoint": knowledge_point, "year": None,
        }
        if not question_text:
            warnings.append(f"第 {block_index} 题未识别到题干")
        if question_type != "judge" and len(options) < 2:
            warnings.append(f"第 {block_index} 题未完整识别选项")
        if not answer:
            warnings.append(f"第 {block_index} 题缺少答案")
        rows.append(row)
    return rows, warnings


def normalize_formatted_rows(rows: list[dict], default_subject: str = "") -> list[dict]:
    """Normalize model output for preview without inventing missing answers."""
    result: list[dict] = []
    seen: set[str] = set()
    for index, raw in enumerate(rows, start=1):
        if not isinstance(raw, dict):
            continue
        row = dict(raw)
        row["question"] = str(row.get("question") or "").strip()
        proposed_id = str(row.get("id") or "").strip() or _stable_formatted_id(row["question"], index)
        if proposed_id in seen:
            proposed_id = f"{proposed_id}_{index}"
        seen.add(proposed_id)
        row["id"] = proposed_id
        row["subject"] = str(row.get("subject") or default_subject or "").strip()
        row["type"] = normalize_question_type(row.get("type"), row.get("answer"))
        row["options"] = _options(row.get("options"))
        row["answer"] = (normalize_fill_answers(row.get("answer")) if row["type"] == "fill"
                         else str(row.get("answer") or "").strip().upper())
        row.setdefault("bank", "")
        row.setdefault("system", "")
        row.setdefault("difficulty", "medium")
        row.setdefault("caseInfo", "")
        row.setdefault("explanation", "")
        points = _knowledge_points(row.get("knowledgePoints") or row.get("knowledgePoint"))
        row["knowledgePoints"] = points[:3]
        row["knowledgePoint"] = "、".join(points[:3])
        row.setdefault("mnemonic", "")
        row.setdefault("year", None)
        result.append(row)
    return result


def recover_trailing_option(row: dict) -> dict:
    """Move one unambiguous, missing Chinese option out of a stem; never invent text."""
    question = copy.deepcopy(row)
    if normalize_question_type(question.get("type"), question.get("answer")) not in {"A1", "A2", "A3", "multiple"}:
        return question
    options = _options(question.get("options"))
    matches = [re.match(r"^([A-Z])[.．、:：)）\s]", option, re.I) for option in options]
    if not 2 <= len(options) < 26 or not all(matches):
        return question
    labels = [match.group(1).upper() for match in matches]
    if labels != sorted(set(labels)):
        return question
    expected = list("ABCDEFGHIJKLMNOPQRSTUVWXYZ"[:len(options) + 1])
    missing = set(expected) - set(labels)
    if len(missing) != 1 or set(labels) - set(expected):
        return question
    label = missing.pop()
    stem = str(question.get("question") or "").strip()
    # Whitespace boundary, one candidate, Chinese text, and no further option markers.
    # ASCII fragments (e.g. vitamin E / D-dimer), multiple missing options and
    # ambiguous inline labels are deliberately left for the user's editor.
    candidates = list(re.finditer(r"\s+" + label + r"[.．、:：)）]?\s*(?=[\u3400-\u9fff])", stem))
    if len(candidates) != 1:
        return question
    match = candidates[0]
    content = stem[match.end():].strip()
    prefix = stem[:match.start()].rstrip()
    if len(prefix) < 4 or not content or len(content) > 500 or re.search(r"\s+[A-Z][.．、:：)）\s]?", content):
        return question
    # A new question/clause after the candidate makes the split uncertain.
    if re.search(r"[\r\n？?。；;]", content):
        return question
    repaired_options = sorted([*options, f"{label}. {content}"], key=lambda item: item[0].upper())
    extensions = question.get("extensions") if isinstance(question.get("extensions"), dict) else {}
    repairs = extensions.get("importRepairs")
    repairs = list(repairs) if isinstance(repairs, list) else []
    repairs.append({"rule": "trailing-option-v1", "label": label, "content": content,
                    "originalQuestion": stem, "originalOptions": options})
    question.update(question=prefix, options=repaired_options,
                    extensions={**extensions, "importRepairs": repairs})
    return question


def normalize_question(row: dict, row_number: int = 0) -> tuple[dict, list[str]]:
    question = recover_trailing_option(row)
    errors: list[str] = []
    for field in REQUIRED_FIELDS:
        value = question.get(field)
        if value is None or not str(value).strip():
            errors.append(f"第 {row_number or '?'} 条缺少字段：{field}")
    question["id"] = str(question.get("id") or "").strip()
    question["subject"] = str(question.get("subject") or "").strip()
    question["question"] = str(question.get("question") or "").strip()
    resolved_type = normalize_question_type(question.get("type"), question.get("answer"))
    question["answer"] = (normalize_fill_answers(question.get("answer")) if resolved_type == "fill"
                          else str(question.get("answer") or "").strip())
    raw_knowledge = question.get("knowledgePoints") or question.get("knowledgePoint")
    raw_knowledge_items = raw_knowledge if isinstance(raw_knowledge, list) else re.split(r"[|,，;；、/]+", str(raw_knowledge or ""))
    raw_knowledge_items = list(dict.fromkeys(str(item).strip() for item in raw_knowledge_items if str(item).strip()))
    knowledge_points = _knowledge_points(raw_knowledge)
    if len(raw_knowledge_items) > 3:
        errors.append(f"第 {row_number or '?'} 条知识点超过 3 个")
    question["knowledgePoints"] = knowledge_points[:3]
    question["knowledgePoint"] = "、".join(knowledge_points[:3])
    question["mnemonic"] = str(question.get("mnemonic") or question.get("记忆口诀") or question.get("口诀") or "").strip()
    raw_tags = normalize_tags(question.get("tags"))
    raw_suggested_tags = normalize_tags(question.get("suggestedTags"))
    if isinstance(question.get("tags"), list) and len(question["tags"]) > 3:
        errors.append(f"第 {row_number or '?'} 条正式标签超过 3 个")
    if isinstance(question.get("suggestedTags"), list) and len(question["suggestedTags"]) > 3:
        errors.append(f"第 {row_number or '?'} 条候选标签超过 3 个")
    question["tags"] = raw_tags
    question["suggestedTags"] = [item for item in raw_suggested_tags if item not in raw_tags]
    question["explanationBlocks"] = normalize_blocks(question.get("explanationBlocks"), question.get("explanation"))
    question["explanationMeta"] = question.get("explanationMeta") if isinstance(question.get("explanationMeta"), dict) else {}
    question["extensions"] = question.get("extensions") if isinstance(question.get("extensions"), dict) else {}
    for issue in question["extensions"].get("importIssues") or []:
        message = str(issue or "").strip()
        if message:
            errors.append(message)
    original_type = str(question.get("type") or "").strip()
    question["type"] = normalize_question_type(original_type, question.get("answer"))
    if original_type and re.sub(r"\s+", "", original_type).casefold() != question["type"].casefold():
        extensions = question.get("extensions") if isinstance(question.get("extensions"), dict) else {}
        extensions.setdefault("originalQuestionType", original_type)
        extensions.setdefault("examQuestionTypeLabel", exam_question_type_label(original_type, question["type"]))
        question["extensions"] = extensions
    if question["type"] not in {"A1", "A2", "A3", "multiple", "judge", "fill"}:
        errors.append(f"第 {row_number or '?'} 条题型无效：{question['type']}")
    difficulty = str(question.get("difficulty") or "medium").strip()
    question["difficulty"] = difficulty
    if difficulty not in {"easy", "medium", "hard"}:
        errors.append(f"第 {row_number or '?'} 条难度无效：{difficulty}")
    question["options"] = _options(question.get("options"))
    if question["type"] == "fill":
        markers = re.findall(r"【(\d+)】", question["question"])
        expected = [str(index) for index in range(1, len(question["answer"]) + 1)]
        if not expected or markers != expected:
            errors.append(f"第 {row_number or '?'} 条填空标记须从【1】连续编号，且与答案空数一致")
        if question["options"]:
            errors.append(f"第 {row_number or '?'} 条填空题不应包含选项")
    elif question["type"] != "judge" and len(question["options"]) < 2:
        errors.append(f"第 {row_number or '?'} 条的非判断题至少需要两个选项")
    if question["type"] in {"A1", "A2", "A3", "multiple"}:
        answer = re.sub(r"[\s,，、;；]+", "", question["answer"]).upper()
        question["answer"] = answer
        allowed = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"[:len(question["options"])]
        if not answer or any(letter not in allowed for letter in answer) or len(set(answer)) != len(answer):
            errors.append(f"第 {row_number or '?'} 条答案必须对应有效且不重复的选项字母")
        elif question["type"] != "multiple" and len(answer) != 1:
            errors.append(f"第 {row_number or '?'} 条单选题只能有一个答案")
        labels = [match.group(1).upper() for option in question["options"]
                  if (match := re.match(r"^([A-Z])[.．、:：)）\s]", option, re.I))]
        empty_labels = [match.group(1).upper() for option in question["options"]
                        if (match := re.fullmatch(r"([A-Z])[.．、:：)）\s]\s*", option, re.I))]
        if empty_labels:
            errors.append(f"第 {row_number or '?'} 条选项内容为空：{'、'.join(empty_labels)}")
        if labels and labels != list(allowed):
            errors.append(f"第 {row_number or '?'} 条选项标记不连续或重复")
            missing = sorted(set("ABCDEFGHIJKLMNOPQRSTUVWXYZ"[:max(len(question["options"]), max(map(ord, labels)) - 64)]) - set(labels))
            if missing:
                errors.append(f"第 {row_number or '?'} 条缺少选项：{'、'.join(missing)}；请补充原文，勿直接重排字母")
        if labels and any(letter not in labels for letter in answer):
            errors.append(f"第 {row_number or '?'} 条答案 {answer} 未对应实际选项标记")
    normalized = normalize_question_v2(question)
    # Internal/legacy automation may still read unknown source keys at top level.
    # v2 export moves the same values into ``extensions`` deterministically.
    known = set(STANDARD_FIELDS) | {"analysis", "evidence", "knowledge", "记忆口诀", "口诀", "解析"}
    normalized.update({key:value for key,value in question.items() if key not in known})
    # Preserve the source string internally; the v2 exporter always regenerates it from blocks.
    normalized["explanation"] = str(question.get("explanation") or "")
    return normalized, errors


def validate_questions(rows: list[dict]) -> list[dict]:
    errors: list[str] = []
    normalized: list[dict] = []
    seen: set[str] = set()
    for index, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            errors.append(f"第 {index} 条不是对象")
            continue
        item, item_errors = normalize_question(row, index)
        errors.extend(item_errors)
        if item["id"] in seen:
            errors.append(f"第 {index} 条 ID 重复：{item['id']}")
        seen.add(item["id"])
        normalized.append(item)
    if not rows:
        errors.append("文件中没有题目")
    if errors:
        raise ImportValidationError(errors)
    return normalized


def _subject_from_chaptered_filename(path: Path) -> str:
    name = path.stem
    name = re.sub(r"[_\-\s]*按章节.*$", "", name, flags=re.I).strip(" _-")
    return name


def _chaptered_options(value: Any) -> list[str]:
    if not isinstance(value, list):
        return _options(value)
    result: list[str] = []
    for index, option in enumerate(value):
        if isinstance(option, dict):
            label = str(option.get("标号") or option.get("label") or chr(65 + index)).strip().upper()
            content = str(option.get("内容") or option.get("content") or "").strip()
            if content:
                result.append(f"{label}. {content}")
        elif str(option).strip():
            result.append(str(option).strip())
    return result


def flatten_chaptered_questions(value: Any, source_path: str | Path) -> list[dict] | None:
    """Adapt [{章节, 题目:[{ID,题型,题干,选项,答案,解析}]}] into standard rows."""
    if not isinstance(value, list) or not value:
        return None
    if not all(isinstance(chapter, dict) and "章节" in chapter and "题目" in chapter for chapter in value):
        return None
    subject = _subject_from_chaptered_filename(Path(source_path))
    rows: list[dict] = []
    for chapter in value:
        system = str(chapter.get("章节") or "").strip()
        questions = chapter.get("题目") or []
        if not isinstance(questions, list):
            continue
        for question in questions:
            if not isinstance(question, dict):
                continue
            answer = str(question.get("答案") or "").strip().upper()
            answer_letters = re.findall(r"[A-E]", answer)
            original_type = str(question.get("题型") or "").strip()
            if original_type in {"判断题", "判断"}:
                question_type = "judge"
            elif len(answer_letters) > 1:
                question_type = "multiple"
            else:
                question_type = "A1"
            rows.append({
                "id": str(question.get("ID") or question.get("id") or "").strip(),
                "bank": "",
                "type": question_type,
                "subject": subject,
                "system": system,
                "difficulty": "medium",
                "caseInfo": "",
                "question": str(question.get("题干") or question.get("question") or "").strip(),
                "options": _chaptered_options(question.get("选项")),
                "answer": "".join(answer_letters) if answer_letters else answer,
                "explanation": str(question.get("解析") or "").strip(),
                "knowledgePoint": str(
                    question.get("知识点") or question.get("knowledgePoint") or question.get("knowledge_point") or ""
                ).strip(),
                "mnemonic": str(question.get("记忆口诀") or question.get("口诀") or question.get("mnemonic") or "").strip(),
                "year": None,
                "sourceQuestionType": original_type,
            })
    return rows


def load_json(path: str | Path, *, validate: bool = True) -> list[dict]:
    source = Path(path)
    with source.open("r", encoding="utf-8-sig") as handle:
        value = json.load(handle)
    return parse_json_questions(value, source, validate=validate)


def parse_json_questions(value: Any, source: str | Path, *, validate: bool = True) -> list[dict]:
    """Shared adapter for a single immutable read of a JSON source."""
    if is_v2(value):
        bank = str((value.get("meta") or {}).get("bank") or "").strip()
        if not bank or any(str(item.get("bank") or "").strip() != bank for item in value["questions"]):
            raise ImportValidationError(["v2 文件只能包含一个题库，且 meta.bank 必须与所有题目一致"])
        return validate_questions(value["questions"]) if validate else value["questions"]
    chaptered = flatten_chaptered_questions(value, source)
    if chaptered is not None:
        return validate_questions(chaptered) if validate else chaptered
    rows = value if isinstance(value, list) else value.get("questions", value.get("data", [])) if isinstance(value, dict) else None
    if not isinstance(rows, list):
        raise ImportValidationError(["JSON 顶层必须是题目数组，或包含 questions/data 数组"])
    return validate_questions(rows) if validate else rows


def excel_headers(path: str | Path) -> list[str]:
    from openpyxl import load_workbook

    v2 = read_xlsx_v2(path)
    if v2 is not None:
        return STANDARD_FIELDS
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    workbook.close()
    return [str(value).strip() if value is not None else "" for value in values]


_PHARM_HEADERS = {"学科", "章节", "题干", "答案", "解析", "A", "B", "C", "D", "E"}


def _repair_mojibake(value: str, codec: str) -> str:
    text = str(value or "")
    pieces = re.split(r"(\s+)", text)
    repaired = []
    for piece in pieces:
        if not piece or piece.isspace():
            repaired.append(piece)
            continue
        try:
            if codec == "latin1-gbk":
                candidate = piece.encode("latin-1").decode("gbk")
            elif codec == "latin1-utf8":
                candidate = piece.encode("latin-1").decode("utf-8")
            else:
                candidate = piece.encode("gbk").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            candidate = piece
        repaired.append(candidate)
    return "".join(repaired)


def _pharmacology_header_codec(headers: list[str]) -> tuple[str, list[str]] | None:
    candidates = [("", headers)]
    for codec in ("latin1-gbk", "latin1-utf8", "gbk-utf8"):
        candidates.append((codec, [_repair_mojibake(item, codec) for item in headers]))
    scored = sorted(candidates, key=lambda item: len(_PHARM_HEADERS & set(item[1])), reverse=True)
    codec, fixed = scored[0]
    best = len(_PHARM_HEADERS & set(fixed))
    original = len(_PHARM_HEADERS & set(headers))
    if best < 8:
        return None
    if codec and best < original + 3:
        return None
    return codec, fixed


def parse_pharmacology_xlsx(
    path: str | Path, progress: Callable[[int, int], None] | None = None,
) -> tuple[list[dict], dict] | None:
    """Adapt the flat pharmacology workbook while retaining invalid rows for preview."""
    from openpyxl import load_workbook

    source = Path(path)
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    first = next(iterator, ())
    headers = [str(value).strip() if value is not None else "" for value in first]
    detected = _pharmacology_header_codec(headers)
    if detected is None:
        workbook.close()
        return None
    codec, fixed_headers = detected
    indexes = {header: index for index, header in enumerate(fixed_headers) if header}
    total = max(0, int(sheet.max_row or 1) - 1)
    rows: list[dict] = []
    digest_counts: dict[str, int] = {}
    digest_positions: dict[str, list[int]] = {}
    missing_explanation = 0
    unsupported_rows = 0

    def clean(value: Any) -> str:
        if value is None:
            return ""
        text = str(value).strip()
        if text.casefold() == "nan":
            return ""
        if codec:
            text = _repair_mojibake(text, codec)
        return unicodedata.normalize("NFC", text)

    def cell(values: tuple, header: str) -> str:
        index = indexes.get(header, -1)
        return clean(values[index]) if 0 <= index < len(values) else ""

    for offset, values in enumerate(iterator, start=2):
        if not any(value not in (None, "") for value in values):
            continue
        subject, system = cell(values, "学科"), cell(values, "章节")
        question, raw_answer = cell(values, "题干"), cell(values, "答案").upper()
        explanation = cell(values, "解析")
        options = [f"{letter}. {cell(values, letter)}" for letter in "ABCDE" if cell(values, letter)]
        answer = "".join(sorted(set(re.findall(r"[A-E]", raw_answer))))
        unsupported = [letter for letter in "FGHIJKLMNOPQRSTUVWXYZ" if cell(values, letter)]
        case_like = bool(
            re.search(r"(?:患者|病人|男性|女性|岁|主诉|入院)", question)
            and re.search(r"(?:天|月|年|疼痛|发热|咳|呕|查体|检查|血|尿|给予|服用)", question)
        )
        qtype = "multiple" if len(answer) > 1 else "A2" if answer and case_like else "A1"
        canonical = json.dumps(
            [subject, system, question, options, answer], ensure_ascii=False,
            separators=(",", ":"), sort_keys=False,
        )
        digest = hashlib.sha1(canonical.encode("utf-8")).hexdigest()[:16]
        occurrence = digest_counts.get(digest, 0) + 1
        digest_counts[digest] = occurrence
        digest_positions.setdefault(digest, []).append(len(rows))
        issues = []
        if not answer:
            issues.append(f"源行 {offset} 缺少答案，程序不会猜测")
        if unsupported:
            issues.append(f"源行 {offset} 的选项 {','.join(unsupported)} 暂不支持，请先调整")
            unsupported_rows += 1
        if not explanation:
            missing_explanation += 1
        rows.append({
            "id": f"xlsx_{digest}_{occurrence}", "bank": "", "questionSource": "",
            "type": qtype, "subject": subject, "system": system, "difficulty": "medium",
            "caseInfo": "", "question": question, "options": options, "answer": answer,
            "explanation": "", "knowledgePoints": [], "tags": [], "suggestedTags": [],
            "extensions": {
                "importedExplanation": explanation,
                "importSource": "pharmacology-xlsx",
                "sourceDocument": source.name,
                "sourceSheet": sheet.title,
                "sourceRow": offset,
                "importIssues": issues,
            },
            "sourceType": "pharmacology-xlsx", "year": None,
        })
        if progress and (len(rows) % 200 == 0 or len(rows) == total):
            progress(len(rows), total)
    workbook.close()
    duplicate_groups = 0
    duplicate_questions = 0
    for digest, positions in digest_positions.items():
        if len(positions) < 2:
            continue
        duplicate_groups += 1
        duplicate_questions += len(positions)
        for position in positions:
            rows[position]["extensions"].setdefault("importWarnings", []).append(
                f"与本文件另外 {len(positions) - 1} 道题内容完全重复，已按要求保留"
            )
    return rows, {
        "subject": next((row["subject"] for row in rows if row["subject"]), ""),
        "rowCount": len(rows), "importableCount": sum(not row["extensions"]["importIssues"] and bool(row["answer"]) for row in rows),
        "missingExplanation": missing_explanation, "duplicateGroups": duplicate_groups,
        "duplicateQuestions": duplicate_questions, "unsupportedOptionRows": unsupported_rows,
        "encodingRepair": codec or "none", "sheet": sheet.title,
    }


def load_excel(path: str | Path, mapping: dict[str, str], *, validate: bool = True) -> list[dict]:
    from openpyxl import load_workbook

    v2 = read_xlsx_v2(path)
    if v2 is not None:
        bank = str((v2.get("meta") or {}).get("bank") or "").strip()
        if not bank or any(str(item.get("bank") or "").strip() != bank for item in v2["questions"]):
            raise ImportValidationError(["v2 文件只能包含一个题库，且 meta.bank 必须与所有题目一致"])
        return validate_questions(v2["questions"]) if validate else v2["questions"]
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        raise ImportValidationError(["Excel 文件为空"])
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    indexes = {header: index for index, header in enumerate(headers) if header}
    result: list[dict] = []
    for values in rows[1:]:
        if not any(value not in (None, "") for value in values):
            continue
        item: dict[str, Any] = {}
        for target, source in mapping.items():
            if source and source in indexes:
                index = indexes[source]
                item[target] = values[index] if index < len(values) else None
        if not item.get("options"):
            option_values = []
            for letter in "ABCDE":
                value = item.pop(f"option{letter}", None)
                if value not in (None, ""):
                    option_values.append(f"{letter}. {str(value).strip()}")
            if option_values:
                item["options"] = option_values
        result.append(item)
    return validate_questions(result) if validate else result


def auto_mapping(headers: list[str]) -> dict[str, str]:
    lower = {header.lower(): header for header in headers if header}
    aliases = {
        "id": ["id", "题目id", "题号"], "subject": ["subject", "学科"],
        "question": ["question", "题目", "题干"], "answer": ["answer", "答案"],
        "options": ["options", "选项"], "type": ["type", "题型"],
        "system": ["system", "章节", "系统"], "caseInfo": ["caseinfo", "病例信息"],
        "difficulty": ["difficulty", "难度"], "knowledgePoint": ["knowledgepoint", "知识点"],
        "knowledgePoints": ["knowledgepoints", "知识点标签"], "mnemonic": ["mnemonic", "记忆口诀", "口诀"],
        "year": ["year", "年份"], "bank": ["bank", "题库"], "explanation": ["explanation", "解析"],
        "tags": ["tags", "题目标签", "标签"], "suggestedTags": ["suggestedtags", "候选标签"],
    }
    for letter in "ABCDE":
        aliases[f"option{letter}"] = [letter, f"{letter}选项", f"选项{letter}", f"option{letter.lower()}"]
    mapping: dict[str, str] = {}
    for target in MAPPING_FIELDS:
        mapping[target] = next((lower[name.lower()] for name in aliases.get(target, [target]) if name.lower() in lower), "")
    return mapping
