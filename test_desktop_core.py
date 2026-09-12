"""Offline tests for the standalone desktop repositories and services."""

from __future__ import annotations

import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

from db_manager import DatabaseManager
from importers import (
    ImportValidationError, auto_mapping, excel_headers, extract_question_text, load_excel, load_json,
    normalize_question, normalize_question_type, parse_common_question_text,
    parse_medical_question_collection_docx, parse_pharmacology_xlsx, question_prompt,
)
from services import (
    ExportService, QuestionFormattingService, RetrievalService, TextbookService,
    canonical_subject, classify_evidence_grade, classify_match_confidence,
    fts_query, normalize_subject_name, retrieval_queries, tokenize_for_fts,
)


SAMPLE = {
    "id": "q_001", "subject": "内科学", "type": "A1", "question": "急性心肌梗死的典型表现是？",
    "options": ["A. 胸痛", "B. 皮疹"], "answer": "A", "customField": "保留",
}


class ImportTests(unittest.TestCase):
    def test_retrieval_query_uses_raw_stem_and_filters_prompt_scaffolding(self):
        question = {
            "prompt_text": "题型：A1 型题 · 单句最佳选择\n有关阿片受体的是？\nA. 干扰项\nC. 纳洛酮",
            "raw": {
                "type": "A1",
                "question": "有关阿片受体的是？",
                "options": ["A. 干扰项", "C. 纳洛酮"],
                "answer": "C",
            },
        }
        rows = retrieval_queries(question)
        self.assertEqual(rows, [("题目核心", "有关阿片受体的是？ C. 纳洛酮")])
        query = fts_query(rows[0][1])
        self.assertIn('"阿片"', query)
        self.assertIn('"受体"', query)
        self.assertIn('"纳洛酮"', query)
        self.assertNotIn('"题型"', query)
        self.assertNotIn('"正确"', query)

    def test_subject_aliases_normalize_for_retrieval_without_rewriting_display_value(self):
        original = "  外基　"
        self.assertEqual(normalize_subject_name(original), "外基")
        self.assertEqual(canonical_subject(original), "外科学")
        self.assertEqual(canonical_subject("外科"), "外科学")
        self.assertEqual(canonical_subject("外科学基础"), "外科学")
        self.assertEqual(
            canonical_subject("诊基", {"retrieval": {"subject_aliases": {"诊基": "诊断学"}}}),
            "诊断学",
        )
        self.assertEqual(original, "  外基　")

    def test_case_queries_and_evidence_grades_follow_local_gates(self):
        question = {"prompt_text": "患者服药后出现皮疹，查体可见红斑，应如何处理？", "raw": {"type": "A2", "question": "患者服药后出现皮疹，查体可见红斑，应如何处理？", "answer": "A", "options": ["A. 停药", "B. 加量"]}}
        self.assertGreaterEqual(len(retrieval_queries(question)), 2)
        strong = {"top_score": .75, "matched_chunks": [{"score": .75, "text": "患者服药后出现皮疹，查体可见红斑时应停药。"}]}
        standard = {"top_score": .60, "matched_chunks": [{"score": .60, "text": "药物不良反应处理原则。"}]}
        low = {"top_score": .49, "matched_chunks": [{"score": .49, "text": "药理学概述。"}]}
        self.assertEqual(classify_evidence_grade(question, strong, .5), "A")
        self.assertEqual(classify_evidence_grade(question, standard, .5), "B")
        self.assertEqual(classify_evidence_grade(question, low, .5), "C")
    def test_json_preserves_unknown_fields(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "questions.json"
            path.write_text(json.dumps([SAMPLE], ensure_ascii=False), encoding="utf-8")
            row = load_json(path)[0]
            self.assertEqual(row["customField"], "保留")

    def test_chaptered_chinese_json_is_flattened(self):
        value = [{
            "章节": "第三章 腹部损伤",
            "题目": [{
                "ID": "chapter_1", "题型": "选择题", "题干": "题干",
                "选项": [{"标号": "A", "内容": "甲"}, {"标号": "B", "内容": "乙"}],
                "答案": "A", "解析": "原解析", "知识点": "腹部损伤处理",
            }],
        }]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "外科学_按章节.json"
            path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
            row = load_json(path)[0]
            self.assertEqual(row["subject"], "外科学")
            self.assertEqual(row["system"], "第三章 腹部损伤")
            self.assertEqual(row["options"], ["A. 甲", "B. 乙"])
            self.assertEqual(row["explanation"], "原解析")
            self.assertEqual(row["knowledgePoint"], "腹部损伤处理")

    def test_duplicate_ids_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "questions.json"
            path.write_text(json.dumps([SAMPLE, SAMPLE], ensure_ascii=False), encoding="utf-8")
            with self.assertRaises(ImportValidationError):
                load_json(path)

    def test_non_judge_requires_options(self):
        _, errors = normalize_question({"id": "1", "subject": "内科", "question": "题干", "answer": "A"}, 1)
        self.assertTrue(any("选项" in item for item in errors))

    def test_excel_separate_option_columns(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "questions.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["题号", "学科", "题干", "答案", "A", "B"])
            sheet.append(["x1", "内科学", "题目", "A", "正确", "错误"])
            workbook.save(path)
            mapping = auto_mapping(excel_headers(path))
            row = load_excel(path, mapping)[0]
            self.assertEqual(row["options"], ["A. 正确", "B. 错误"])

    def test_pharmacology_xlsx_preserves_duplicates_special_text_and_stable_ids(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "药理学题库.xlsx"
            workbook = Workbook(); sheet = workbook.active
            sheet.append(["学科", "章节", "题干", "答案", "解析", "A", "B", "C", "D", "E"])
            row = ["药理学", "药动学", "t₁/₂ 与 Na⁺ 浓度的关系是？", "A", "原解析", "甲", "乙", "丙", "丁", "戊"]
            sheet.append(row); sheet.append(row)
            sheet.append(["药理学", "病例", "患者服药后出现不适，查体异常，应选？", "", "", "甲", "乙", "丙", "丁", "戊"])
            workbook.save(path)
            first_rows, stats = parse_pharmacology_xlsx(path)
            second_rows, _stats = parse_pharmacology_xlsx(path)
        self.assertEqual(stats["rowCount"], 3)
        self.assertEqual(stats["importableCount"], 2)
        self.assertEqual(stats["duplicateGroups"], 1)
        self.assertEqual(first_rows[0]["question"], "t₁/₂ 与 Na⁺ 浓度的关系是？")
        self.assertEqual([row["id"] for row in first_rows], [row["id"] for row in second_rows])
        self.assertNotEqual(first_rows[0]["id"], first_rows[1]["id"])
        self.assertEqual(first_rows[0]["extensions"]["importedExplanation"], "原解析")
        self.assertEqual(first_rows[0]["explanation"], "")
        self.assertTrue(normalize_question(first_rows[2], 3)[1])

    def test_pharmacology_xlsx_repairs_only_a_matching_mojibake_header(self):
        from openpyxl import Workbook
        garble = lambda value: value.encode("utf-8").decode("latin-1")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "garbled.xlsx"
            workbook = Workbook(); sheet = workbook.active
            headers = ["学科", "章节", "题干", "答案", "解析", "A", "B", "C", "D", "E"]
            sheet.append([garble(value) if len(value) > 1 else value for value in headers])
            sheet.append([garble("药理学"), garble("总论"), garble("题干"), "A", "", garble("甲"), garble("乙"), "", "", ""])
            workbook.save(path)
            rows, stats = parse_pharmacology_xlsx(path)
        self.assertEqual(stats["encodingRepair"], "latin1-utf8")
        self.assertEqual(rows[0]["subject"], "药理学")
        self.assertEqual(rows[0]["question"], "题干")

    def test_common_raw_text_is_organized_locally(self):
        raw = """1. 急性心肌梗死最常见的症状是？
A. 胸痛
B. 皮疹
C. 耳鸣
D. 腹泻
答案：A

2. 下列哪些属于多选？
A. 甲
B. 乙
C. 丙
答案：AB"""
        rows, warnings = parse_common_question_text(raw, "内科学")
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["answer"], "A")
        self.assertEqual(rows[1]["type"], "multiple")
        self.assertEqual(rows[0]["subject"], "内科学")
        self.assertEqual(warnings, [])

    def test_common_raw_text_reads_knowledge_point_marker(self):
        raw = """1. 急性心肌梗死最常见的症状是？
A. 胸痛
B. 皮疹
答案：A
知识点：急性心肌梗死的临床表现
解析：持续胸痛是典型表现。"""
        rows, warnings = parse_common_question_text(raw, "内科学")
        self.assertEqual(warnings, [])
        self.assertEqual(rows[0]["knowledgePoint"], "急性心肌梗死的临床表现")
        self.assertNotIn("知识点", rows[0]["explanation"])

    def test_knowledge_points_are_capped_and_exported_as_tags(self):
        normalized = normalize_question({**SAMPLE, "knowledgePoint": "休克、液体复苏、损伤控制"}, 1)[0]
        self.assertEqual(normalized["knowledgePoints"], ["休克", "液体复苏", "损伤控制"])
        _, errors = normalize_question({**SAMPLE, "knowledgePoint": "标签甲、标签乙、标签丙、标签丁"}, 1)
        self.assertTrue(any("超过 3 个" in item for item in errors))

    def test_missing_answer_is_flagged_not_invented(self):
        rows, warnings = parse_common_question_text("1. 题干？\nA. 甲\nB. 乙", "内科学")
        self.assertEqual(rows[0]["answer"], "")
        self.assertTrue(any("缺少答案" in item for item in warnings))

    def test_formatter_uses_ai_fallback_and_keeps_empty_answer(self):
        class FakeLLM:
            def organize_questions(self, _text, subject):
                return [{"question": "无法由本地规则识别的题目", "subject": subject, "options": ["A. 甲", "B. 乙"], "answer": ""}]

        result = QuestionFormattingService(FakeLLM()).organize("无编号的杂乱内容", "诊断学", True)
        self.assertEqual(result["method"], "deepseek")
        self.assertEqual(result["rows"][0]["answer"], "")
        self.assertTrue(result["rows"][0]["id"].startswith("formatted_"))
        self.assertTrue(result["warnings"][0])

    def test_docx_raw_text_is_read(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "questions.docx"
            document = Document()
            document.add_paragraph("1. 题干？")
            document.add_paragraph("A. 甲")
            document.add_paragraph("B. 乙")
            document.add_paragraph("答案：A")
            document.save(path)
            self.assertIn("答案：A", extract_question_text(path))

    def test_styled_medical_collection_docx_preserves_chapters_and_exam_types(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "病理学_题目集合.docx"
            document = Document()
            document.add_paragraph("病理学 - 题目集合", style="Title")
            document.add_paragraph("目录概览", style="Heading 1")
            document.add_paragraph("第一章 绪论 (2题)", style="List Bullet")
            document.add_paragraph("第一章 绪论", style="Heading 1")
            for number, raw_type, answer in ((1, "X型题", "AC"), (2, "B型题", "B")):
                document.add_paragraph(f"第一章 绪论 - 第{number}题", style="Heading 2")
                document.add_paragraph(f"题型：{raw_type} | 题号：7 | 文件：2026-01-01-00-00-0{number}")
                document.add_paragraph("题目：", style="Heading 3")
                document.add_paragraph(f"测试题干 {number}")
                document.add_paragraph("选项：", style="Heading 3")
                for option in ("A. 甲", "B. 乙", "C. 丙", "D. 丁", "E. 戊"):
                    document.add_paragraph(option, style="List Bullet")
                document.add_paragraph("正确答案：", style="Heading 3")
                document.add_paragraph(answer)
                document.add_paragraph("解析：", style="Heading 3")
                document.add_paragraph("原题解析。")
                document.add_paragraph("考点讨论：", style="Heading 3")
                document.add_paragraph("考点讨论内容。")
            document.save(path)
            rows, stats = parse_medical_question_collection_docx(path)
            self.assertEqual(stats["questionCount"], 2)
            self.assertEqual(stats["chapterCount"], 1)
            self.assertEqual(rows[0]["type"], "multiple")
            self.assertEqual(rows[1]["type"], "A1")
            self.assertEqual(rows[1]["extensions"]["examQuestionTypeLabel"], "B 型题 · 配伍题")
            self.assertTrue(rows[1]["extensions"]["questionGroupId"].startswith("ykb_b_"))
            self.assertEqual(rows[0]["system"], "第一章 绪论")
            self.assertEqual(len(rows[0]["explanationBlocks"]), 2)
            self.assertEqual(stats["issues"], [])

    def test_question_type_aliases_are_normalized_and_prompted(self):
        self.assertEqual(normalize_question_type("A1型题", "A"), "A1")
        self.assertEqual(normalize_question_type("单选题", "A"), "A1")
        self.assertEqual(normalize_question_type("X型题", "AC"), "multiple")
        self.assertEqual(normalize_question_type("B型题", "B"), "A1")
        self.assertEqual(normalize_question_type("A3/A4型题", "D"), "A3")
        prompt = question_prompt({
            "type":"A1", "answer":"B", "question":"题干", "options":["A. 甲","B. 乙"],
            "extensions":{"originalQuestionType":"B型题"},
        })
        self.assertIn("题型：B 型题 · 配伍题", prompt)

    def test_structured_docx_organizer_does_not_call_ai(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "内科学_题目集合.docx"
            document = Document()
            document.add_paragraph("内科学 - 题目集合", style="Title")
            document.add_paragraph("第一章 循环", style="Heading 1")
            document.add_paragraph("第一章 循环 - 第1题", style="Heading 2")
            document.add_paragraph("题型：A2型题 | 题号：1 | 文件：marker-1")
            for title, values in (("题目：", ["患者胸痛，最可能的诊断是"]), ("选项：", ["A. 甲", "B. 乙"]), ("正确答案：", ["A"])):
                document.add_paragraph(title, style="Heading 3")
                for value in values: document.add_paragraph(value)
            document.save(path)
            result = QuestionFormattingService().organize_file(path)
            self.assertEqual(result["method"], "structured-docx")
            self.assertEqual(result["rows"][0]["type"], "A2")
            self.assertEqual(result["warnings"], [[]])

    def test_llm_formatter_accepts_json_code_fence(self):
        from llm_client import LLMClient
        client = object.__new__(LLMClient)
        client._complete = lambda _prompt: '```json\n[{"question":"题干","options":["A.甲","B.乙"],"answer":"A"}]\n```'
        rows = client.organize_questions("原文", "内科学")
        self.assertEqual(rows[0]["answer"], "A")


class RepositoryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.database = DatabaseManager(str(Path(self.temp.name) / "app.db"))

    def tearDown(self):
        self.database.close()
        self.temp.cleanup()

    def test_question_sets_isolate_same_external_id(self):
        first = self.database.create_question_set("第一批", "one.json", "json", [SAMPLE])
        second = self.database.create_question_set("第二批", "two.json", "json", [SAMPLE])
        self.assertNotEqual(first, second)
        self.assertEqual(len(self.database.list_imported_questions(first)), 1)
        self.assertEqual(len(self.database.list_imported_questions(second)), 1)

    def test_partial_import_appends_to_same_set(self):
        first = dict(SAMPLE)
        second = dict(SAMPLE)
        second["id"] = "q_002"
        set_id = self.database.create_question_set("分批导入", "pasted://manual", "formatted", [first])
        self.database.append_questions_to_set(set_id, [second])
        rows = self.database.list_imported_questions(set_id)
        self.assertEqual(len(rows), 2)
        question_set = next(item for item in self.database.list_question_sets() if item["id"] == set_id)
        self.assertEqual(question_set["question_count"], 2)

    def test_bulk_subject_and_selected_queue(self):
        first = dict(SAMPLE)
        second = dict(SAMPLE)
        second["id"] = "q_bulk_2"
        set_id = self.database.create_question_set("批量操作", "source.json", "json", [first, second])
        rows = self.database.list_imported_questions(set_id)
        first_id, second_id = rows[0]["id"], rows[1]["id"]
        updated = self.database.bulk_update_question_subject([first_id, second_id], "外科学")
        self.assertEqual(updated, 2)
        self.assertEqual(self.database.get_imported_question(first_id)["raw"]["subject"], "外科学")
        self.database.bulk_queue_questions([first_id, second_id])
        selected = self.database.queued_questions(set_id, [second_id])
        self.assertEqual([row["id"] for row in selected], [second_id])

    def test_cross_page_filter_returns_only_lightweight_matching_ids(self):
        questions = [
            {**SAMPLE, "id": "bulk_001", "question": "病史采集第一题"},
            {**SAMPLE, "id": "bulk_002", "question": "心电图第二题"},
            {**SAMPLE, "id": "bulk_003", "question": "病史采集第三题"},
        ]
        set_id = self.database.create_question_set("跨页选择", "bulk.json", "json", questions)
        all_ids = self.database.list_question_ids_by_filter(set_id)
        matching_ids = self.database.list_question_ids_by_filter(set_id, pipeline_status="queued", search="病史采集")
        self.assertEqual(len(all_ids), 3)
        self.assertEqual(matching_ids, [all_ids[0], all_ids[2]])

    def test_candidate_tag_needs_three_questions_and_can_be_deleted_everywhere(self):
        questions = [{**SAMPLE, "id": f"tag_{index}"} for index in range(1, 4)]
        set_id = self.database.create_question_set("标签治理", "tags.json", "json", questions)
        rows = self.database.list_imported_questions(set_id)

        first = self.database.get_imported_question(rows[0]["id"])["raw"]
        first["suggestedTags"] = ["病史问诊"]
        self.database.conn.execute(
            "UPDATE imported_questions SET raw_json=? WHERE id=?",
            (json.dumps(first, ensure_ascii=False), rows[0]["id"]),
        )
        self.database.conn.commit()
        candidate = self.database.list_question_tags(set_id, status="candidate")[0]
        self.assertEqual(candidate["usage_count"], 1)
        with self.assertRaisesRegex(ValueError, "3 道题"):
            self.database.update_question_tag(candidate["id"], status="active")

        for row in rows[1:]:
            raw = self.database.get_imported_question(row["id"])["raw"]
            raw["suggestedTags"] = ["病史问诊"]
            self.database.conn.execute(
                "UPDATE imported_questions SET raw_json=? WHERE id=?",
                (json.dumps(raw, ensure_ascii=False), row["id"]),
            )
        self.database.conn.commit()
        active = self.database.list_question_tags(set_id, status="active")[0]
        self.assertEqual(active["usage_count"], 3)
        result = self.database.delete_question_tag(active["id"])
        self.assertEqual(result, {"label": "病史问诊", "changed": 3})
        self.assertEqual(self.database.list_question_tags(set_id), [])
        for row in rows:
            raw = self.database.get_imported_question(row["id"])["raw"]
            self.assertNotIn("病史问诊", raw.get("tags", []))
            self.assertNotIn("病史问诊", raw.get("suggestedTags", []))

    def test_review_requires_opening_question(self):
        set_id = self.database.create_question_set("测试", "one.json", "json", [SAMPLE])
        question = self.database.list_imported_questions(set_id)[0]
        self.database.save_generated_result(question["id"], "generated", "textbook", 0.8, "解析", [])
        with self.assertRaises(ValueError):
            self.database.review_question(question["id"], "approved", "解析")
        self.database.mark_question_viewed(question["id"])
        self.database.review_question(question["id"], "approved", "解析")
        self.assertEqual(self.database.get_imported_question(question["id"])["review_status"], "approved")

    def test_invalid_generated_output_is_reclassified_as_generation_error(self):
        questions = [{**SAMPLE, "id": "missing_source"}, {**SAMPLE, "id": "general_ok"}]
        set_id = self.database.create_question_set("状态修复", "repair.json", "json", questions)
        rows = self.database.list_imported_questions(set_id)
        self.database.save_generated_result(
            rows[0]["id"], "generated", "textbook", 0.81,
            '{"explanationBlocks": [{"type": "paragraph"}]}', [],
        )
        general_package={"explanationBlocks":[{"section":"answerBasis","type":"paragraph","text":"简要说明"},
                                                        {"section":"pitfalls","type":"table","columns":["易错点","辨析"],"rows":[["其他项","不符合"]]}],
                         "briefExplanation":"简要说明正确答案成立的核心原因。","knowledgePoints":[],
                         "tags":["通识标签"],"suggestedTags":[],"explanationMeta":{"mode":"general_knowledge","evidence":[]}}
        self.database.save_generated_result(
            rows[1]["id"], "generated", "general_knowledge", 0.0,
            "未在所选教材中找到直接依据。简要说明。", [], package=general_package,
        )

        self.assertEqual(self.database.reclassify_invalid_generated_questions(set_id), 1)
        failed = self.database.get_imported_question(rows[0]["id"])
        self.assertEqual(failed["pipeline_status"], "error")
        self.assertEqual(failed["retrieval_status"], "error")
        self.assertEqual(failed["generation_status"], "validation_error")
        self.assertEqual(self.database.get_imported_question(rows[1]["id"])["pipeline_status"], "generated")

    def test_tag_backfill_preserves_approved_explanation_and_evidence(self):
        set_id = self.database.create_question_set("标签补齐", "tags.json", "json", [SAMPLE])
        question = self.database.list_imported_questions(set_id)[0]
        blocks = [
            {"section":"analysis","type":"paragraph","title":"考点解析","text":"已有考点"},
            {"section":"answerBasis","type":"paragraph","title":"答案依据","text":"已有依据"},
            {"section":"pitfalls","type":"table","title":"易错点","columns":["选项","辨析"],"rows":[["B","不符合"]]},
        ]
        package = {"explanationBlocks":blocks,"briefExplanation":"已有简析","knowledgePoints":[],
                   "tags":[],"suggestedTags":[],"explanationMeta":{"mode":"textbook","evidence":[]}}
        evidence = [{"textbook":"内科学","sourcePage":12,"quote":"教材原文","score":.9}]
        self.database.save_generated_result(
            question["id"], "generated", "textbook", .9, "原解析", evidence, package=package,
        )
        self.database.mark_question_viewed(question["id"])
        before = self.database.get_imported_question(question["id"])
        self.database.review_question(question["id"], "approved", before["explanation"])
        approved_before = self.database.get_imported_question(question["id"])
        candidates = self.database.list_missing_tag_questions(set_id, include_approved=True)
        self.assertEqual(len(candidates), 1)
        self.assertTrue(candidates[0]["repairable"])
        self.database.apply_tag_backfill(question["id"], ["急性冠脉综合征"])
        after = self.database.get_imported_question(question["id"])
        self.assertEqual(after["review_status"], "approved")
        self.assertEqual(after["explanation"], approved_before["explanation"])
        self.assertEqual(after["evidence"], evidence)
        self.assertEqual(after["raw"]["tags"], [])
        self.assertEqual(after["raw"]["suggestedTags"], ["急性冠脉综合征"])

    def test_generation_failure_does_not_overwrite_existing_content(self):
        set_id = self.database.create_question_set("失败保留", "failure.json", "json", [SAMPLE])
        question = self.database.list_imported_questions(set_id)[0]
        evidence = [{"textbook":"内科学","sourcePage":1,"quote":"旧证据","score":.8}]
        self.database.save_generated_result(
            question["id"], "generated", "textbook", .8, "原有解析", evidence,
        )
        self.database.save_generation_failure(
            question["id"], "error", "模型输出格式修复失败", score=.7,
        )
        after = self.database.get_imported_question(question["id"])
        self.assertEqual(after["explanation"], "原有解析")
        self.assertEqual(after["evidence"], evidence)
        self.assertEqual(after["pipeline_status"], "error")
        self.assertEqual(after["retrieval_status"], "matched")
        self.assertEqual(after["generation_status"], "validation_error")

    def test_high_score_legacy_unmatched_is_reclassified_without_losing_evidence(self):
        questions = [{**SAMPLE, "id": "legacy_high"}, {**SAMPLE, "id": "legacy_low"}]
        set_id = self.database.create_question_set("legacy", "legacy.json", "json", questions)
        rows = self.database.list_imported_questions(set_id)
        evidence = [
            {"textbook": "教材甲", "sourcePage": 70, "quote": "证据甲", "score": 0.748},
            {"textbook": "教材乙", "sourcePage": 12, "quote": "证据乙", "score": 0.701},
        ]
        self.database.save_generated_result(
            rows[0]["id"], "unmatched", "", 0.748, "", evidence, "模型输出格式修复失败",
        )
        self.database.save_generated_result(
            rows[1]["id"], "unmatched", "", 0.49, "", evidence, "相似度未达到阈值",
        )

        self.assertEqual(self.database.reconcile_high_score_unmatched(0.6), 1)
        high = self.database.get_imported_question(rows[0]["id"])
        low = self.database.get_imported_question(rows[1]["id"])
        self.assertEqual(
            (high["pipeline_status"], high["retrieval_status"], high["generation_status"]),
            ("error", "matched", "validation_error"),
        )
        self.assertEqual(high["match_confidence"], "strong")
        self.assertEqual(high["evidence"], evidence)
        self.assertEqual(low["pipeline_status"], "unmatched")

        retry_rows = self.database.matched_generation_failures(set_id)
        self.assertEqual([item["id"] for item in retry_rows], [rows[0]["id"]])
        self.assertEqual(
            self.database.bulk_queue_questions([rows[0]["id"]], preserve_retrieval=True), 1,
        )
        queued = self.database.get_imported_question(rows[0]["id"])
        self.assertEqual(queued["pipeline_status"], "queued")
        self.assertEqual(queued["retrieval_status"], "matched")
        self.assertEqual(queued["evidence"], evidence)

    def test_pending_textbook_output_below_new_threshold_leaves_review_queue(self):
        questions = [
            {**SAMPLE, "id": "pending_textbook"},
            {**SAMPLE, "id": "approved_textbook"},
            {**SAMPLE, "id": "general_result"},
        ]
        set_id = self.database.create_question_set("阈值重分类", "threshold.json", "json", questions)
        rows = self.database.list_imported_questions(set_id)
        evidence = [{"textbook": "内科学", "sourcePage": 12, "quote": "教材原文", "score": 0.55}]
        for row in rows[:2]:
            self.database.save_generated_result(
                row["id"], "generated", "textbook", 0.55, "教材解析", evidence,
            )
        self.database.mark_question_viewed(rows[1]["id"])
        self.database.review_question(rows[1]["id"], "approved", "教材解析")
        self.database.save_generated_result(
            rows[2]["id"], "generated", "general_knowledge", 0.40, "通识解析", [],
        )

        self.assertEqual(
            self.database.reclassify_below_threshold_generated_questions(0.60, set_id), 1,
        )
        pending = self.database.get_imported_question(rows[0]["id"])
        approved = self.database.get_imported_question(rows[1]["id"])
        general = self.database.get_imported_question(rows[2]["id"])
        self.assertEqual(
            (pending["pipeline_status"], pending["retrieval_status"], pending["generation_status"]),
            ("unmatched", "unmatched", "pending"),
        )
        self.assertEqual(approved["pipeline_status"], "generated")
        self.assertEqual(general["pipeline_status"], "generated")

    def test_match_confidence_uses_top_score_and_consensus(self):
        self.assertEqual(classify_match_confidence([0.748], 0.6), "strong")
        self.assertEqual(classify_match_confidence([0.66, 0.65], 0.6), "strong")
        self.assertEqual(classify_match_confidence([0.62, 0.51], 0.6), "standard")
        self.assertEqual(classify_match_confidence([0.59, 0.58], 0.6), "low")

    def test_quick_review_only_approves_strict_textbook_candidates(self):
        questions = []
        for suffix in ("good", "low", "warning"):
            questions.append({**SAMPLE, "id": f"quick_{suffix}"})
        set_id = self.database.create_question_set("快速审核", "quick.json", "json", questions)
        rows = self.database.list_imported_questions(set_id)
        evidence = [{
            "textbook": "内科学", "source_file": "内科学.pdf", "source_page": 88,
            "pdf_page": 103, "score": 0.88,
            "text": "教材明确记载急性心肌梗死的典型表现为持续性胸痛，并可伴随大汗。",
        }]
        blocks = [
            {"section": "analysis", "type": "paragraph", "title": "考点解析", "text": "本题考查急性心肌梗死的典型临床表现，需要结合持续胸痛的特点识别。"},
            {"section": "answerBasis", "type": "paragraph", "title": "正确答案依据", "text": "教材指出持续性胸痛是典型表现，因此胸痛选项符合教材证据。"},
            {"section": "pitfalls", "type": "table", "title": "易错点提示", "columns": ["易错点", "正确辨析"], "rows": [["皮疹", "不属于急性心肌梗死的典型表现"]]},
        ]
        base_package = {
            "explanationBlocks": blocks, "knowledgePoints": ["急性心肌梗死典型表现为持续胸痛"],
            "briefExplanation": "持续胸痛符合急性心肌梗死的典型临床表现。",
            "tags": ["心肌梗死表现"], "suggestedTags": [],
            "explanationMeta": {"mode": "textbook", "evidence": evidence},
        }
        self.database.save_generated_result(rows[0]["id"], "generated", "textbook", 0.88, "", evidence, package=base_package)
        self.database.save_generated_result(rows[1]["id"], "generated", "textbook", 0.65, "", evidence, package=base_package)
        warning_package = json.loads(json.dumps(base_package, ensure_ascii=False))
        warning_package["explanationMeta"]["reviewWarnings"] = ["结构修复失败"]
        self.database.save_generated_result(rows[2]["id"], "generated", "textbook", 0.90, "", evidence, package=warning_package)

        preview = self.database.quick_review_candidates(set_id, 0.75)
        self.assertEqual([item["external_id"] for item in preview["eligible"]], ["quick_good"])
        self.assertEqual(preview["excluded"]["相似度低于门槛"], 1)
        self.assertEqual(preview["excluded"]["存在格式警告"], 1)
        result = self.database.bulk_approve_quick_review(
            set_id, 0.75, [rows[0]["id"], rows[1]["id"]],
        )
        self.assertEqual(result, {"approved": 1, "skipped": 1})
        self.assertEqual(self.database.get_imported_question(rows[0]["id"])["review_status"], "approved")
        self.assertEqual(self.database.get_imported_question(rows[1]["id"])["review_status"], "pending")
        action = self.database.conn.execute(
            "SELECT note FROM review_actions WHERE question_pk=? ORDER BY id DESC", (rows[0]["id"],)
        ).fetchone()
        self.assertIn("快速审核", action["note"])

    def test_knowledge_point_update_preserves_pipeline_and_exports(self):
        set_id = self.database.create_question_set("知识点", "one.json", "json", [SAMPLE])
        question = self.database.list_imported_questions(set_id)[0]
        self.database.save_generated_result(question["id"], "generated", "textbook", 0.8, "解析", [])
        self.database.update_question_knowledge_point(question["id"], "冠状动脉粥样硬化性心脏病")
        updated = self.database.get_imported_question(question["id"])
        self.assertEqual(updated["raw"]["knowledgePoint"], "冠状动脉粥样硬化性心脏病")
        self.assertEqual(updated["pipeline_status"], "generated")
        self.database.mark_question_viewed(question["id"])
        self.database.review_question(question["id"], "approved", "解析")
        paths = ExportService(self.database).export(set_id, self.temp.name)
        exported = json.loads(Path(paths["json"]).read_text(encoding="utf-8"))
        self.assertEqual(exported["questions"][0]["knowledgePoints"], ["冠状动脉粥样硬化性心脏病"])

    def test_fts_and_embedding_rerank(self):
        root = Path(self.temp.name) / "books"
        root.mkdir()
        pdf = root / "internal.pdf"
        pdf.write_bytes(b"pdf")
        library_id = self.database.add_library("内科学教材", "内科学", str(pdf))
        file_id = self.database.upsert_file_metadata(library_id, str(pdf.resolve()), pdf.name, 3, 1, "hash")
        self.database.replace_file_content(
            file_id, library_id,
            [{"page_number": 1, "text": "急性心肌梗死胸痛", "extraction_method": "text"}],
            [{"page_number": 1, "chunk_index": 0, "chunk_text": "急性心肌梗死常表现为持续胸痛",
              "search_text": " ".join(tokenize_for_fts("急性心肌梗死常表现为持续胸痛")),
              "embedding": [1.0, 0.0], "extraction_method": "text"}],
        )
        config = {"retrieval": {"bm25_top_k": 20, "final_top_k": 3, "similarity_threshold": 0.5}}
        result = RetrievalService(self.database, config).retrieve("内科学", "心肌梗死胸痛", [1.0, 0.0])
        self.assertTrue(result["is_matched"])
        self.assertEqual(result["matched_chunks"][0]["source_page"], 1)

    def test_selected_textbooks_are_isolated_and_diversified(self):
        library_ids = []
        books = (
            ("内科学甲版", "内科学"),
            ("病理学教材", "病理学"),
            ("未选教材", "外科学"),
        )
        for index, (name, subject) in enumerate(books, start=1):
            pdf = Path(self.temp.name) / f"selected-{index}.pdf"
            pdf.write_bytes(b"pdf")
            library_id = self.database.add_library(name, subject, str(pdf), f"第{index}版")
            library_ids.append(library_id)
            file_id = self.database.upsert_file_metadata(
                library_id, str(pdf.resolve()), pdf.name, 3, index, f"selected-{index}"
            )
            self.database.replace_file_content(
                file_id, library_id,
                [{"page_number": index, "text": "心肌梗死胸痛", "extraction_method": "text"}],
                [{"page_number": index, "chunk_index": 0, "chunk_text": f"{name}记载心肌梗死胸痛",
                  "search_text": " ".join(tokenize_for_fts("心肌梗死胸痛")),
                  "embedding": [1.0, 0.0], "extraction_method": "text"}],
            )
        config = {"retrieval": {"bm25_top_k": 20, "final_top_k": 3, "similarity_threshold": 0.5}}
        result = RetrievalService(self.database, config).retrieve(
            "内科学", "心肌梗死胸痛", [1.0, 0.0], library_ids[:2],
        )
        sources = {item["library_id"] for item in result["matched_chunks"]}
        self.assertEqual(sources, set(library_ids[:2]))
        self.assertNotIn(library_ids[2], sources)

    def test_textbook_page_calibration_keeps_pdf_locator(self):
        pdf = Path(self.temp.name) / "surgery.pdf"
        pdf.write_bytes(b"pdf")
        library_id = self.database.add_library(
            "外科学教材", "外科学", str(pdf), "第十版", page_offset=39
        )
        file_id = self.database.upsert_file_metadata(
            library_id, str(pdf.resolve()), pdf.name, 3, 1, "hash-page-calibration"
        )
        self.database.replace_file_content(
            file_id, library_id,
            [{"page_number": 527, "text": "急性胰腺炎腹痛", "extraction_method": "text"}],
            [{
                "page_number": 527, "chunk_index": 0,
                "chunk_text": "急性胰腺炎主要表现为腹痛",
                "search_text": " ".join(tokenize_for_fts("急性胰腺炎主要表现为腹痛")),
                "embedding": [1.0, 0.0], "extraction_method": "text",
            }],
        )
        config = {"retrieval": {"bm25_top_k": 20, "final_top_k": 3, "similarity_threshold": 0.5}}
        result = RetrievalService(self.database, config).retrieve(
            "外科学", "急性胰腺炎腹痛", [1.0, 0.0]
        )
        evidence = result["matched_chunks"][0]
        self.assertEqual(evidence["source_page"], 488)
        self.assertEqual(evidence["pdf_page"], 527)

    def test_printed_page_offset_is_inferred_from_page_headers(self):
        pdf = Path(self.temp.name) / "auto-pages.pdf"
        pdf.write_bytes(b"pdf")
        library_id = self.database.add_library("外科学", "外科学", str(pdf))
        file_id = self.database.upsert_file_metadata(
            library_id, str(pdf.resolve()), pdf.name, 3, 1, "auto-page-hash"
        )
        pages = []
        chunks = []
        for pdf_page in (527, 547, 567):
            printed = pdf_page - 39
            text = f"{printed}\n第七篇 腹部外科疾病\n正文"
            pages.append({"page_number": pdf_page, "text": text, "extraction_method": "text"})
            chunks.append({
                "page_number": pdf_page, "chunk_index": 0, "chunk_text": text,
                "search_text": " ".join(tokenize_for_fts(text)),
                "embedding": [1.0, 0.0], "extraction_method": "text",
            })
        self.database.replace_file_content(file_id, library_id, pages, chunks)
        set_id = self.database.create_question_set("页码迁移", "source.json", "json", [SAMPLE])
        question = self.database.list_imported_questions(set_id)[0]
        self.database.save_generated_result(
            question["id"], "generated", "textbook", 0.9,
            "教材出处：课本第547页",
            [{"source_path": str(pdf.resolve()), "source_file": pdf.name, "source_page": 547}],
        )
        self.assertEqual(self.database.infer_library_page_offset(library_id), 39)
        self.assertEqual(self.database.get_library(library_id)["page_offset"], 39)
        migrated = self.database.get_imported_question(question["id"])
        self.assertEqual(migrated["evidence"][0]["source_page"], 508)
        self.assertEqual(migrated["evidence"][0]["pdf_page"], 547)
        self.assertIn("课本第508页", migrated["explanation"])

    def test_delete_questions_updates_set_count(self):
        first = dict(SAMPLE)
        second = dict(SAMPLE)
        second["id"] = "delete_2"
        set_id = self.database.create_question_set("删除测试", "source.json", "json", [first, second])
        rows = self.database.list_imported_questions(set_id)
        self.assertEqual(self.database.delete_imported_questions([rows[0]["id"]]), 1)
        question_set = next(row for row in self.database.list_question_sets() if row["id"] == set_id)
        self.assertEqual(question_set["question_count"], 1)
        self.assertEqual(self.database.trash_count(), 0)

    def test_delete_question_set_is_permanent_and_keeps_textbooks(self):
        pdf = Path(self.temp.name) / "keep.pdf"
        pdf.write_bytes(b"%PDF-1.4\n")
        library_id = self.database.add_library("保留教材", "内科学", str(pdf))
        set_id = self.database.create_question_set("整库删除", "source.json", "json", [SAMPLE])
        job_id = self.database.create_job("generate", "旧任务", {"set_id": set_id})
        self.database.update_job(job_id, status="completed")
        result = self.database.delete_question_set(set_id)
        self.assertTrue(result["deleted"])
        self.assertEqual(result["questions"], 1)
        self.assertFalse(self.database.list_question_sets())
        self.assertIsNone(self.database.get_job(job_id))
        self.assertEqual(self.database.get_library(library_id)["name"], "保留教材")

    def test_active_job_blocks_question_set_deletion(self):
        set_id = self.database.create_question_set("运行中题库", "source.json", "json", [SAMPLE])
        question = self.database.list_imported_questions(set_id)[0]
        self.database.create_job("general", "运行中", {"question_pk": question["id"]})
        with self.assertRaises(ValueError):
            self.database.delete_question_set(set_id)
        self.assertEqual(len(self.database.list_imported_questions(set_id)), 1)

    def test_job_records_can_be_permanently_deleted(self):
        finished = self.database.create_job("generate", "完成", {})
        active = self.database.create_job("generate", "等待", {})
        self.database.update_job(finished, status="completed")
        result = self.database.delete_jobs([finished, active])
        self.assertEqual(result["deleted"], 1)
        self.assertEqual(result["blocked"], [active])
        self.assertIsNone(self.database.get_job(finished))
        self.assertIsNotNone(self.database.get_job(active))

    def test_clear_question_data_preserves_libraries_and_jobs(self):
        pdf = Path(self.temp.name) / "keep-index.pdf"
        pdf.write_bytes(b"%PDF-1.4\n")
        library_id = self.database.add_library("保留索引", "内科学", str(pdf))
        self.database.create_question_set("待清空", "source.json", "json", [SAMPLE])
        job_id = self.database.create_job("generate", "历史任务", {})
        self.database.update_job(job_id, status="completed")
        counts = self.database.clear_all_question_data()
        self.assertEqual(counts["questions"], 1)
        self.assertFalse(self.database.list_question_sets())
        self.assertEqual(self.database.get_library(library_id)["name"], "保留索引")
        self.assertIsNotNone(self.database.get_job(job_id))

    def test_existing_database_gets_page_offset_column(self):
        self.database.close()
        path = Path(self.temp.name) / "old-library.db"
        connection = sqlite3.connect(path)
        connection.execute(
            "CREATE TABLE libraries (id INTEGER PRIMARY KEY, name TEXT, subject TEXT, "
            "version TEXT, root_path TEXT, active INTEGER, last_scanned_at TEXT, created_at TEXT)"
        )
        connection.commit()
        connection.close()
        with DatabaseManager(str(path)) as migrated:
            columns = {row["name"] for row in migrated.conn.execute("PRAGMA table_info(libraries)")}
            self.assertIn("page_offset", columns)

    def test_export_only_approved_and_preserves_source(self):
        set_id = self.database.create_question_set("测试", "one.json", "json", [SAMPLE])
        question = self.database.list_imported_questions(set_id)[0]
        evidence = [{"source_file": "内科学.pdf", "source_page": 5, "score": 0.9}]
        self.database.save_generated_result(question["id"], "generated", "textbook", 0.9, "审核解析", evidence)
        self.database.mark_question_viewed(question["id"])
        self.database.review_question(question["id"], "approved", "审核解析")
        paths = ExportService(self.database).export(set_id, self.temp.name)
        exported = json.loads(Path(paths["json"]).read_text(encoding="utf-8"))
        self.assertEqual(exported["schema"], "medlearning.question-set")
        item = exported["questions"][0]
        self.assertEqual(item["extensions"]["customField"], "保留")
        self.assertIn("审核解析", item["explanation"])
        self.assertEqual(item["explanationMeta"]["evidence"][0]["sourcePage"], 5)

    def test_job_recovery(self):
        job_id = self.database.create_job("generate", "测试任务", {"set_id": "x"})
        self.database.update_job(job_id, status="running")
        self.database.recover_incomplete_jobs()
        self.assertEqual(self.database.get_job(job_id)["status"], "queued")

    def test_single_pdf_index_and_incomplete_retry(self):
        import fitz

        class FakeEmbedding:
            def get_embeddings(self, texts):
                return [[1.0, 0.0] for _ in texts]

        pdf = Path(self.temp.name) / "book.pdf"
        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 72), "Clinical medicine textbook content for indexing and retrieval.")
        document.save(pdf)
        document.close()
        library_id = self.database.add_library("教材", "内科学", str(pdf))
        config = {
            "retrieval": {"chunk_size": 500, "chunk_overlap": 50},
            "memory_guard": {"max_memory_mb": 2048, "batch_process_size": 50},
        }
        service = TextbookService(self.database, config, FakeEmbedding())
        first = service.scan_library(library_id)
        self.assertEqual(first["changed"], 1)
        file_row = self.database.get_file_by_path(str(pdf.resolve()))
        self.assertGreater(self.database.file_chunk_count(file_row["id"]), 0)
        self.database.conn.execute("UPDATE textbook_files SET status='indexing' WHERE id=?", (file_row["id"],))
        self.database.conn.execute("DELETE FROM chunks_v2 WHERE file_id=?", (file_row["id"],))
        self.database.conn.commit()
        second = service.scan_library(library_id)
        self.assertEqual(second["changed"], 1)
        self.assertGreater(self.database.file_chunk_count(file_row["id"]), 0)

    def test_empty_pdf_marks_index_failed(self):
        import fitz

        class FakeEmbedding:
            def get_embeddings(self, texts):
                return [[1.0] for _ in texts]

        pdf = Path(self.temp.name) / "empty.pdf"
        document = fitz.open()
        document.new_page()
        document.save(pdf)
        document.close()
        library_id = self.database.add_library("空教材", "内科学", str(pdf))
        config = {
            "retrieval": {"chunk_size": 500, "chunk_overlap": 50},
            "memory_guard": {"max_memory_mb": 2048, "batch_process_size": 50},
        }
        with self.assertRaises(RuntimeError):
            TextbookService(self.database, config, FakeEmbedding()).scan_library(library_id)
        row = self.database.get_file_by_path(str(pdf.resolve()))
        self.assertEqual(row["status"], "error")


class MigrationTests(unittest.TestCase):
    def test_v1_questions_are_copied_without_deleting_source(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "legacy.db"
            connection = sqlite3.connect(path)
            connection.executescript(
                """
                CREATE TABLE questions (
                    id INTEGER PRIMARY KEY, question_id TEXT UNIQUE, question_text TEXT,
                    matched_chunks TEXT, match_score REAL, match_status TEXT,
                    explanation TEXT, error_message TEXT, updated_at TEXT
                );
                CREATE TABLE textbook_chunks (
                    id INTEGER PRIMARY KEY, source_file TEXT, source_page INTEGER,
                    chunk_index INTEGER, chunk_text TEXT, embedding BLOB, embedding_dim INTEGER
                );
                INSERT INTO questions VALUES
                    (1, 'old_1', '旧题目', '[]', 0.8, 'matched', '旧解析', '', '2026-01-01');
                """
            )
            connection.commit()
            connection.close()
            with DatabaseManager(str(path)) as database:
                question_sets = database.list_question_sets()
                self.assertEqual(question_sets[0]["name"], "旧版导入")
                migrated = database.list_imported_questions(question_sets[0]["id"])[0]
                self.assertEqual(database.get_imported_question(migrated["id"])["explanation"], "旧解析")
                self.assertEqual(database.conn.execute("SELECT COUNT(*) FROM questions").fetchone()[0], 1)


class GlobalQuestionQueryTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.database = DatabaseManager(str(Path(self.tmp.name) / "test.db"))
        self.set_a = self.database.create_question_set("集A", "a.json", "json", [
            {**SAMPLE, "id": "a1", "subject": "外科学", "questionSource": "医考帮",
             "tags": ["休克"], "question": "休克的首选处理？"},
            {**SAMPLE, "id": "a2", "subject": "内科学", "questionSource": "本校历年题",
             "tags": ["心衰"], "question": "心衰的诱因？"},
        ])
        self.set_b = self.database.create_question_set("集B", "b.json", "json", [
            {**SAMPLE, "id": "b1", "subject": "外科学", "questionSource": "医考帮",
             "tags": ["休克", "围手术期"], "question": "围手术期补液原则？"},
        ])

    def tearDown(self) -> None:
        self.database.close()
        self.tmp.cleanup()

    def test_global_list_spans_sets(self) -> None:
        rows = self.database.list_questions_global()
        self.assertEqual([row["external_id"] for row in rows], ["a1", "a2", "b1"])
        self.assertEqual(self.database.count_questions_global(), 3)
        self.assertEqual(rows[0]["tags"], ["休克"])
        self.assertEqual(rows[0]["question_source"], "医考帮")

    def test_global_filters(self) -> None:
        self.assertEqual(
            [r["external_id"] for r in self.database.list_questions_global(set_id=self.set_b)], ["b1"])
        self.assertEqual(
            [r["external_id"] for r in self.database.list_questions_global(subject="外科学")], ["a1", "b1"])
        self.assertEqual(
            [r["external_id"] for r in self.database.list_questions_global(tag="休克")], ["a1", "b1"])
        self.assertEqual(
            [r["external_id"] for r in self.database.list_questions_global(question_source="本校历年题")], ["a2"])
        self.assertEqual(
            [r["external_id"] for r in self.database.list_questions_global(search="补液")], ["b1"])
        self.assertEqual(self.database.count_questions_global(subject="外科学", tag="休克"), 2)
        self.assertEqual(self.database.list_question_ids_global(subject="内科学"),
                         [self.database.list_questions_global(subject="内科学")[0]["id"]])

    def test_global_status_filters(self) -> None:
        rows = self.database.list_questions_global(review_status="pending", pipeline_status="queued")
        self.assertEqual(len(rows), 3)
        self.assertEqual(self.database.list_questions_global(pipeline_status="matched_generation_error"), [])

    def test_global_pagination(self) -> None:
        page = self.database.list_questions_global(limit=2, offset=2)
        self.assertEqual([row["external_id"] for row in page], ["b1"])

    def test_filter_option_lists(self) -> None:
        self.assertEqual(self.database.all_question_subjects(), ["内科学", "外科学"])
        self.assertEqual(self.database.all_question_sources(), ["医考帮", "本校历年题"])
        # 三题中 "休克" 出现 2 次（<3，不晋升 active），目录无 active 标签
        self.assertEqual(self.database.all_tag_labels(), [])

    def test_count_trash_with_search(self) -> None:
        first = self.database.list_questions_global()[0]["id"]
        self.database.move_questions_to_trash([first])
        self.assertEqual(self.database.count_trash(), 1)
        self.assertEqual(self.database.count_trash(search="a1"), 1)
        self.assertEqual(self.database.count_trash(search="不存在"), 0)

    def test_restore_trash_preserves_v8_columns(self) -> None:
        target = self.database.list_questions_global(search="心衰")[0]
        self.database.conn.execute(
            "UPDATE imported_questions SET retrieval_status='matched', generation_status='generated',"
            " match_confidence='strong' WHERE id=?", (target["id"],))
        self.database.conn.commit()
        self.database.move_questions_to_trash([target["id"]])
        trash_id = self.database.list_trash()[0]["id"]
        result = self.database.restore_trash([trash_id])
        self.assertEqual(result, {"restored": 1, "conflicts": []})
        restored = self.database.list_questions_global(search="心衰")[0]
        self.assertEqual(restored["retrieval_status"], "matched")
        self.assertEqual(restored["generation_status"], "generated")
        self.assertEqual(restored["match_confidence"], "strong")


if __name__ == "__main__":
    unittest.main()
